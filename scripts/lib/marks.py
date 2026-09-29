"""Pinned annotation colors, mark wording, and label vocabularies.

grader.md / grading-checker.md require these to be identical in every graded
file, so they live in one place and every writer (annotate.py,
mark_review.py) and reader (audit_graded.py) uses the same values.
"""
from __future__ import annotations

GRADER_RED = "FF0000"
CHECKER_BLUE = "0000FF"
BAD_FILL = "FFC7CE"   # Excel "Bad" cell style: light red fill ...
BAD_FONT = "9C0006"   # ... with dark red font

GRADING_COMPLETED = "Grading Completed"
REVIEW_PASSED = "Review Passed"
REVIEW_FAILED = "Review FAILED"
REVIEW_MARKS = (REVIEW_PASSED, REVIEW_FAILED)
SUMMARY_SHEET = "Grading Summary"

# Labels a grader annotation may carry ("partial" is annotated as INCOMPLETE).
VERDICT_LABELS = ("INCORRECT", "INCOMPLETE")

DISCREPANCY_TYPES = (
    "verdict_mismatch",
    "label_mismatch",
    "outdated_rubric",
    "missed_question",
    "annotation_placement",
    "explanation_error",
    "incomplete_coverage",
    "cell_highlight_missing",
    "missing_keypoint_not_flagged",
    "format_error",
)


# The graded file's provenance record, written by annotate.py as custom
# document properties (docProps/custom.xml — invisible on the page, kept by
# Word and Excel): which solutions file and which verified rubric the file was
# graded against, and which rubric question each annotation belongs to. The
# JSON is split into <=250-character properties META_NAME.1, .2, … (Office's
# limit for a text property). Each annotation also carries a hidden tag — a
# bookmark (docx) or defined name (xlsx) named ANNOTATION_TAG+n.
META_NAME = "grading-assistant"
META_CHUNK = 250
ANNOTATION_TAG = "ga_ann_"
CUSTOM_NS = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
VT_NS = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
FMTID = "{D5CDD505-2E9C-101B-9397-08002B2CF9AE}"


def _chunks(meta: dict) -> list[tuple[str, str]]:
    import json
    raw = json.dumps(meta, ensure_ascii=False, separators=(",", ":"))
    return [(f"{META_NAME}.{i + 1}", raw[i * META_CHUNK:(i + 1) * META_CHUNK])
            for i in range((len(raw) + META_CHUNK - 1) // META_CHUNK)]


def write_meta_docx(document, meta: dict) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from lxml import etree

    package = document.part.package
    part = next((r.target_part for r in package.rels.values() if r.reltype == RT.CUSTOM_PROPERTIES), None)
    if part is None:
        root = etree.Element(f"{{{CUSTOM_NS}}}Properties", nsmap={None: CUSTOM_NS, "vt": VT_NS})
    else:
        root = etree.fromstring(part.blob)
    for prop in [p for p in root if (p.get("name") or "").startswith(META_NAME + ".")]:
        root.remove(prop)
    pid = max([int(p.get("pid", 1)) for p in root] + [1])
    for name, value in _chunks(meta):
        pid += 1
        prop = etree.SubElement(root, f"{{{CUSTOM_NS}}}property",
                                fmtid=FMTID, pid=str(pid), name=name)
        etree.SubElement(prop, f"{{{VT_NS}}}lpwstr").text = value
    blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    if part is None:
        part = Part(PackURI("/docProps/custom.xml"),
                    "application/vnd.openxmlformats-officedocument.custom-properties+xml",
                    blob, package)
        package.relate_to(part, RT.CUSTOM_PROPERTIES)
    else:
        part._blob = blob


def write_meta_xlsx(workbook, meta: dict) -> None:
    from openpyxl.packaging.custom import StringProperty
    props = workbook.custom_doc_props
    for name in [p.name for p in props if p.name.startswith(META_NAME + ".")]:
        del props[name]
    for name, value in _chunks(meta):
        props.append(StringProperty(name=name, value=value))


def read_meta(graded_file) -> dict | None:
    """The provenance record of a graded .docx/.xlsx, or None."""
    import json
    import zipfile
    from lxml import etree

    try:
        with zipfile.ZipFile(graded_file) as zf:
            if "docProps/custom.xml" not in zf.namelist():
                return None
            root = etree.fromstring(zf.read("docProps/custom.xml"))
    except (OSError, zipfile.BadZipFile, etree.XMLSyntaxError):
        return None
    parts = {}
    for prop in root:
        name = prop.get("name") or ""
        if name.startswith(META_NAME + ".") and name.rsplit(".", 1)[1].isdigit():
            parts[int(name.rsplit(".", 1)[1])] = "".join(prop.itertext())
    if not parts:
        return None
    try:
        return json.loads("".join(parts[i] for i in sorted(parts)))
    except json.JSONDecodeError:
        return None


def mark_state(graded_file) -> dict:
    """Which pinned marks a graded .docx / .xlsx already carries (exact
    text): {"grading_completed": bool, "reviewed": bool}."""
    from pathlib import Path

    path = Path(graded_file)
    if path.suffix.lower() == ".docx":
        import docx
        return mark_state_of_document(docx.Document(str(path)))
    import openpyxl
    wb = openpyxl.load_workbook(str(path))
    if SUMMARY_SHEET not in wb.sheetnames:
        return {"grading_completed": False, "reviewed": False}
    ws = wb[SUMMARY_SHEET]
    return {"grading_completed": ws["A1"].value == GRADING_COMPLETED,
            "reviewed": any(v in REVIEW_MARKS for row in ws.iter_rows(values_only=True) for v in row)}


def mark_state_of_document(document) -> dict:
    texts = [p.text.strip() for p in document.paragraphs]
    return {"grading_completed": GRADING_COMPLETED in texts,
            "reviewed": any(t in REVIEW_MARKS for t in texts)}


def rgb_hex(color) -> str | None:
    """Normalize a python-docx RGBColor / openpyxl Color to 'RRGGBB' (or None)."""
    if color is None:
        return None
    rgb = getattr(color, "rgb", color)
    if rgb is None or not isinstance(rgb, str):
        # python-docx RGBColor stringifies to 'RRGGBB'; openpyxl theme/indexed
        # colors have no literal rgb string.
        try:
            rgb = str(rgb)
        except Exception:
            return None
    rgb = rgb.upper()
    if len(rgb) == 8:  # openpyxl ARGB
        rgb = rgb[2:]
    return rgb if len(rgb) == 6 else None
