#!/usr/bin/env python3
"""Mechanical audit of a `_Graded.docx|xlsx` for the checker — in two steps
that make blind grading structural rather than an honour rule.

1. Status (default): is the `Grading Completed` mark there, is the file
   already reviewed, and was it graded against the current verified rubric?
   Nothing about the grader's verdicts is shown or written.

2. Report (`--report --blind <checker_blind.json>`): the checker first
   commits its own blind verdict for EVERY rubric question to a file; only
   then does this produce the report (audit.json, in the checker's own
   folder), which records the blind file's hash. It contains:
   - every grader annotation — its rubric question (from the hidden tag
     annotate.py writes), verdict label, full text, and the id of the
     paragraph/cell it sits under (the extraction view's ids);
   - `comparison`: per question, the checker's blind verdict vs the
     grader's (no annotation = correct): agree / verdict_mismatch (correct vs
     not) / label_mismatch (INCORRECT vs INCOMPLETE);
   - `problems`: facts, not judgments — the mark not exact, original student
     content altered or deleted, wrong colour shade, annotation text over a
     non-empty cell, a highlight without an annotation (or vice versa), an
     annotation that is not the closest empty visible cell, a missing
     question tag, and grading against an outdated rubric.
   mark_review.py refuses to write the final mark without this report.

Whether a verdict or explanation is *right* stays the checker's own work.

Usage:
    .venv/Scripts/python scripts/audit_graded.py <graded> <submission> --solutions <solutions>
    .venv/Scripts/python scripts/audit_graded.py <graded> <submission> --solutions <solutions> --report --blind <checker_blind.json>

checker_blind.json:
    {"verdicts": [{"question": "Q1", "verdict": "correct"},
                  {"question": "Q2", "verdict": "incorrect", "explanation": "…"},
                  {"question": "Q3", "verdict": "incomplete", "explanation": "…"}]}
  verdict: correct | incorrect | incomplete | unreadable — one per rubric question.
"""
from __future__ import annotations

import argparse
import difflib
import hashlib
import json
import re
import sys
import warnings
from pathlib import Path

import _venv  # noqa: F401  — must precede third-party imports (re-runs under .venv)

import docx
import openpyxl
from docx.text.paragraph import Paragraph
from openpyxl.chartsheet import Chartsheet
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib import cache, marks  # noqa: E402
import annotate  # noqa: E402
import extract  # noqa: E402
import rubric  # noqa: E402
from extract_docx import body_blocks, element_text, has_image  # noqa: E402

LABEL_RE = re.compile(r"^\s*(INCORRECT|INCOMPLETE|PARTIAL|CORRECT|UNREADABLE)\b\s*:?", re.I)
# Near-miss variants of the pinned marks ("Grading Completed ✓", "Grading
# Completed | Review Passed", "review failed") — still marks, but not exact.
COMPLETED_RE = re.compile(r"\bgrading completed\b", re.I)
REVIEW_RE = re.compile(r"^\W*review (passed|failed)\b", re.I)


def short(text, limit: int = 160) -> str:
    text = "" if text is None else str(text)
    return text if len(text) <= limit else text[:limit - 1] + "…"


def one_line(text: str) -> str:
    return " / ".join(part for part in text.split("\n") if part.strip())


def verdict_label(text: str) -> str | None:
    m = LABEL_RE.match(text or "")
    return m.group(1).upper() if m else None


# ------------------------------------------------------------------ documents

def body_items(document) -> list[dict]:
    """Top-level body content in order, with the extraction view's ids."""
    items = []
    for bid, kind, obj in body_blocks(document):
        if kind == "t":
            items.append({"kind": "t", "id": bid, "table": obj, "key": "T:" + bid,
                          "text": one_line(element_text(obj._tbl)[0])})
            continue
        el = obj._p if kind == "p" else obj
        text = element_text(el)[0]
        items.append({"kind": kind, "id": bid, "para": obj if kind == "p" else None,
                      "text": text, "key": f"{kind.upper()}:{text}",
                      "blank": not text.strip() and not has_image(el)})
    return items


def run_colors(para: Paragraph) -> list[str | None]:
    return [marks.rgb_hex(r.font.color.rgb) if r.font.color is not None and r.font.color.type else None
            for r in para.runs if r.text.strip()]


def classify(para: Paragraph, text: str) -> dict:
    colors = run_colors(para)
    bolds = [bool(r.bold) for r in para.runs if r.text.strip()]
    color_set = sorted({c or "none" for c in colors})
    stripped = text.strip()
    if stripped == marks.GRADING_COMPLETED or (len(stripped) < 60 and COMPLETED_RE.search(stripped)):
        kind = "grading_completed_mark"
    elif stripped in marks.REVIEW_MARKS or (len(stripped) < 60 and REVIEW_RE.match(stripped)):
        kind = "review_mark"
    elif color_set == [marks.GRADER_RED]:
        kind = "grader_annotation"
    elif color_set == [marks.CHECKER_BLUE]:
        kind = "checker_text"
    else:
        kind = "unexpected_insert"
    tags = para._p.xpath("./w:bookmarkStart/@w:name")
    tag = next((int(t[len(marks.ANNOTATION_TAG):]) for t in tags
                if t.startswith(marks.ANNOTATION_TAG) and t[len(marks.ANNOTATION_TAG):].isdigit()), None)
    return {"kind": kind, "text": text, "colors": color_set, "all_bold": bool(bolds) and all(bolds),
            "tag": tag}


def audit_cell_paragraphs(src_cell, g_cell, anchor: str, report: dict) -> None:
    src = [element_text(p._p)[0] for p in src_cell.paragraphs]
    g_paras = g_cell.paragraphs
    g = [element_text(p._p)[0] for p in g_paras]
    sm = difflib.SequenceMatcher(a=src, b=g, autojunk=False)
    for op, a0, a1, b0, b1 in sm.get_opcodes():
        if op in ("replace", "delete"):
            report["problems"].append({"type": "format_error", "where": anchor,
                                       "detail": "original table-cell text altered or removed",
                                       "original": [short(t) for t in src[a0:a1]],
                                       "graded": [short(t) for t in g[b0:b1]]})
        if op in ("insert", "replace"):
            for j in range(b0, b1):
                info = classify(g_paras[j], g[j])
                if op == "replace" and info["kind"] != "grader_annotation":
                    continue
                report["_inserted"].append({**info, "after": anchor,
                                            "answer_text": short(one_line(element_text(src_cell._tc)[0]))})


def audit_document(graded: Path, submission: Path) -> dict:
    report = {"type": "docx", "graded_file": str(graded), "problems": [], "_inserted": []}
    src_items = body_items(docx.Document(str(extract.source_docx_for(submission))))
    g_doc = docx.Document(str(graded))
    g_items = body_items(g_doc)
    sm = difflib.SequenceMatcher(a=[i["key"] for i in src_items], b=[i["key"] for i in g_items],
                                 autojunk=False)
    anchor_of_graded: dict[int, dict] = {}
    for op, a0, a1, b0, b1 in sm.get_opcodes():
        if op == "equal":
            for k in range(a1 - a0):
                anchor_of_graded[b0 + k] = src_items[a0 + k]
        else:
            if op in ("replace", "delete"):
                report["problems"].append({
                    "type": "format_error",
                    "detail": "original content altered or removed",
                    "original": [f"{i['id']}: {short(i['text'])}" for i in src_items[a0:a1]],
                    "graded": [short(i["text"]) for i in g_items[b0:b1]]})
            for j in range(b0, b1):
                item = g_items[j]
                if item["kind"] != "p":  # annotate.py only ever inserts paragraphs
                    continue
                info = classify(item["para"], item["text"])
                if op == "replace" and info["kind"] != "grader_annotation":
                    continue
                prev = next((anchor_of_graded[k] for k in range(j - 1, -1, -1) if k in anchor_of_graded),
                            None)
                report["_inserted"].append({
                    **info, "_pos": j,
                    "after": prev["id"] if prev else None,
                    "answer_text": (short(prev["text"]) or "(image)") if prev else None,
                    "anchor_is_blank": bool(prev and prev["kind"] == "p" and prev["blank"]),
                })

    # Annotations placed inside table cells.
    src_tables = [i["table"] for i in src_items if i["kind"] == "t"]
    g_tables = [i["table"] for i in g_items if i["kind"] == "t"]
    if len(src_tables) != len(g_tables):
        report["problems"].append({"type": "format_error",
                                   "detail": f"table count changed {len(src_tables)} -> {len(g_tables)}"})
    for t, (st, gt) in enumerate(zip(src_tables, g_tables)):
        seen = set()
        for r, (srow, grow) in enumerate(zip(st.rows, gt.rows)):
            for c, (sc, gc) in enumerate(zip(srow.cells, grow.cells)):
                if id(gc._tc) in seen:
                    continue
                seen.add(id(gc._tc))
                audit_cell_paragraphs(sc, gc, f"t{t}:r{r}c{c}", report)

    inserted = report.pop("_inserted")
    annotations = [i for i in inserted if i["kind"] == "grader_annotation"]
    mark_items = [i for i in inserted if i["kind"] == "grading_completed_mark"]
    exact = [i for i in mark_items if i["text"] == marks.GRADING_COMPLETED]
    mark = {"present": bool(mark_items), "count": len(mark_items),
            "exact_text": bool(exact) and len(mark_items) == 1,
            "color_ok": bool(mark_items) and all(i["colors"] == [marks.GRADER_RED] for i in mark_items),
            "bold": bool(mark_items) and all(i["all_bold"] for i in mark_items)}
    if mark_items:
        mark_pos = mark_items[-1].get("_pos", 10 ** 9)
        mark["placed_after_all_annotations"] = all(a.get("_pos", -1) < mark_pos for a in annotations)
    report["grading_completed"] = mark
    report["review_mark_present"] = any(i["kind"] == "review_mark" for i in inserted)

    for n, a in enumerate(annotations, 1):
        if a.get("anchor_is_blank"):
            report["problems"].append({"type": "annotation_placement", "annotation": n,
                                       "detail": f"sits below a blank paragraph ({a['after']}) — "
                                                 "check it is immediately below the answer"})
        if not verdict_label(a["text"]):
            report["problems"].append({"type": "format_error", "annotation": n,
                                       "detail": "annotation does not start with a verdict label "
                                                 "(INCORRECT:/INCOMPLETE:)"})
    for i in inserted:
        if i["kind"] == "unexpected_insert":
            report["problems"].append({"type": "format_error",
                                       "detail": f"inserted text with colors {i['colors']} "
                                                 f"(grader red must be {marks.GRADER_RED})",
                                       "after": i.get("after"), "text": short(i["text"])})
    if mark_items and not (mark["exact_text"] and mark["color_ok"] and mark["bold"]):
        report["problems"].append({"type": "format_error",
                                   "detail": f"'{marks.GRADING_COMPLETED}' mark is not exact "
                                             f"(needs exact text, {marks.GRADER_RED}, bold, once)"})
    if mark_items and not mark.get("placed_after_all_annotations", True):
        report["problems"].append({"type": "annotation_placement",
                                   "detail": "an annotation appears after the Grading Completed mark"})

    report["annotations"] = [{"n": n, "tag": a["tag"], "after": a["after"],
                              "answer_text": a["answer_text"],
                              "verdict_label": verdict_label(a["text"]), "text": a["text"]}
                             for n, a in enumerate(annotations, 1)]
    return report


# ------------------------------------------------------------------ workbooks

def _fill_hex(cell) -> str | None:
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return None
    return marks.rgb_hex(fill.fgColor)


def _font_hex(cell) -> str | None:
    return marks.rgb_hex(cell.font.color) if cell.font is not None and cell.font.color is not None else None


def _comparable(value):
    """A cell value comparable across two loads: openpyxl returns array and
    data-table formulas as objects without equality."""
    if isinstance(value, (ArrayFormula, DataTableFormula)):
        return (type(value).__name__, getattr(value, "text", None), tuple(value))
    return value


def audit_workbook(graded: Path, submission: Path) -> dict:
    from openpyxl.cell.cell import MergedCell
    from openpyxl.utils.cell import get_column_letter

    report = {"type": "xlsx", "graded_file": str(graded), "problems": []}
    source, notes = annotate._workbook_source(submission)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        orig = openpyxl.load_workbook(str(source)) if source else annotate._rebuild_from_xls(submission)
        g = openpyxl.load_workbook(str(graded))
    if notes:
        report["notes"] = notes

    extra = [s for s in g.sheetnames if s not in orig.sheetnames and s != marks.SUMMARY_SHEET]
    missing = [s for s in orig.sheetnames if s not in g.sheetnames]
    if extra or missing:
        report["problems"].append({"type": "format_error",
                                   "detail": f"sheets added {extra} / removed {missing}"})

    annotations = []
    for name in orig.sheetnames:
        if name not in g.sheetnames:
            continue
        ows, gws = orig[name], g[name]
        if isinstance(ows, Chartsheet) or isinstance(gws, Chartsheet):
            continue  # a chart on its own tab has no cells to annotate or alter
        max_row = max(ows.max_row, gws.max_row)
        max_col = max(ows.max_column, gws.max_column)
        highlighted, notes_cells = [], {}
        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                oc, gc = ows.cell(row=r, column=c), gws.cell(row=r, column=c)
                ov, gv = _comparable(oc.value), _comparable(gc.value)
                coord = f"{get_column_letter(c)}{r}"
                o_empty = ov in (None, "")
                if not o_empty and gv != ov:
                    report["problems"].append({
                        "type": "format_error", "where": f"{name}!{coord}",
                        "detail": "original cell value/formula changed (an annotation must never "
                                  "overwrite a non-empty cell)",
                        "original": short(ov), "graded": short(gv)})
                if o_empty and gv not in (None, ""):
                    font = _font_hex(gc)
                    if font != marks.GRADER_RED:
                        report["problems"].append({"type": "format_error", "where": f"{name}!{coord}",
                                                   "detail": f"inserted text is {font}, not {marks.GRADER_RED}",
                                                   "text": short(gv)})
                    notes_cells[(r, c)] = gv
                g_fill, o_fill = _fill_hex(gc), _fill_hex(oc)
                if g_fill == marks.BAD_FILL and o_fill != marks.BAD_FILL:
                    if _font_hex(gc) != marks.BAD_FONT:
                        report["problems"].append({"type": "cell_highlight_missing",
                                                   "where": f"{name}!{coord}",
                                                   "detail": f"highlight font is {_font_hex(gc)}, not {marks.BAD_FONT}"})
                    highlighted.append((r, c))
                elif g_fill != o_fill and g_fill is not None:
                    report["problems"].append({"type": "cell_highlight_missing", "where": f"{name}!{coord}",
                                               "detail": f"fill changed to {g_fill} (highlight must be "
                                                         f"{marks.BAD_FILL})"})

        # Pair highlighted answers with annotations by replaying annotate.py's
        # placement over the ORIGINAL sheet: answers in row order, each
        # taking the first cell of its search order that was empty and not
        # already taken by an answer or an earlier annotation.
        highlight_set = set(highlighted)
        hidden_rows, hidden_cols = annotate.hidden_rows_cols(ows)
        for (r, c) in notes_cells:
            if r in hidden_rows or c in hidden_cols:
                report["problems"].append({"type": "annotation_placement",
                                           "where": f"{name}!{get_column_letter(c)}{r}",
                                           "detail": "annotation is in a hidden row/column — "
                                                     "the student will never see it"})

        def blocked(rc) -> bool:
            oc = ows.cell(row=rc[0], column=rc[1])
            return (oc.value not in (None, "") or isinstance(oc, MergedCell)
                    or rc in highlight_set or rc[0] in hidden_rows or rc[1] in hidden_cols)

        claimed: dict[tuple, tuple] = {}  # note cell -> answer cell
        for (r, c) in sorted(highlighted):
            for rc in annotate.search_order(r, c):
                if blocked(rc) or rc in claimed:
                    continue
                if rc in notes_cells:
                    claimed[rc] = (r, c)
                else:
                    report["problems"].append({"type": "cell_highlight_missing",
                                               "where": f"{name}!{get_column_letter(c)}{r}",
                                               "detail": "highlighted cell has no annotation in the "
                                                         "closest empty cell"})
                break
        # Extra notes on one answer stack further along its search order.
        for rc in sorted(set(notes_cells) - set(claimed)):
            best = None
            for (r, c) in highlighted:
                for rank, cand in enumerate(annotate.search_order(r, c)):
                    if cand == rc:
                        if best is None or rank < best[0]:
                            best = (rank, (r, c))
                        break
                    if not (blocked(cand) or cand in notes_cells):
                        break
            if best:
                claimed[rc] = best[1]
        unpaired_notes = {rc: t for rc, t in notes_cells.items() if rc not in claimed}
        for rc, (r, c) in sorted(claimed.items(), key=lambda kv: (kv[1], kv[0])):
            text = notes_cells[rc]
            annotations.append({"sheet": name, "answer_cell": f"{get_column_letter(c)}{r}",
                                "answer_value": short(gws.cell(row=r, column=c).value, 80),
                                "annotation_cell": f"{get_column_letter(rc[1])}{rc[0]}",
                                "verdict_label": verdict_label(str(text)), "text": str(text)})
        for (r, c), text in unpaired_notes.items():
            annotations.append({"sheet": name, "answer_cell": None,
                                "annotation_cell": f"{get_column_letter(c)}{r}",
                                "verdict_label": verdict_label(str(text)), "text": str(text)})
            report["problems"].append({"type": "annotation_placement",
                                       "where": f"{name}!{get_column_letter(c)}{r}",
                                       "detail": "annotation is not the closest empty cell to any "
                                                 "highlighted answer cell (or its answer is not highlighted)"})

    summary = {"present": marks.SUMMARY_SHEET in g.sheetnames}
    review_present = False
    if summary["present"]:
        ws = g[marks.SUMMARY_SHEET]
        a1 = ws["A1"]
        summary.update({"is_last_tab": g.sheetnames[-1] == marks.SUMMARY_SHEET,
                        "exact_text": a1.value == marks.GRADING_COMPLETED,
                        "color_ok": _font_hex(a1) == marks.GRADER_RED,
                        "bold": bool(a1.font and a1.font.bold)})
        review_present = any(v in marks.REVIEW_MARKS for row in ws.iter_rows(values_only=True) for v in row)
        if not all(summary[k] for k in ("is_last_tab", "exact_text", "color_ok", "bold")):
            report["problems"].append({"type": "format_error",
                                       "detail": f"'{marks.SUMMARY_SHEET}' tab / A1 mark not exact: {summary}"})
    report["grading_completed"] = summary
    report["review_mark_present"] = review_present
    tag_of_cell = {}
    for tag_name, defined in g.defined_names.items():
        if tag_name.startswith(marks.ANNOTATION_TAG) and tag_name[len(marks.ANNOTATION_TAG):].isdigit():
            for sheet_name, ref in defined.destinations:
                # openpyxl leaves a quoted name's doubled apostrophes (O''Brien) as is
                tag_of_cell[(sheet_name.replace("''", "'"), ref.replace("$", ""))] = int(tag_name[len(marks.ANNOTATION_TAG):])
    for n, a in enumerate(annotations, 1):
        a["n"] = n
        a["tag"] = tag_of_cell.get((a["sheet"], a["annotation_cell"]))
        if not a["verdict_label"]:
            report["problems"].append({"type": "format_error", "annotation": n,
                                       "detail": "annotation does not start with a verdict label"})
    report["annotations"] = annotations
    return report


BLIND_VERDICTS = ("correct", "incorrect", "incomplete", "unreadable")


def provenance_problems(graded: Path, solutions: Path) -> tuple[dict | None, list[dict]]:
    """The graded file's record vs the current verified rubric."""
    meta = marks.read_meta(graded)
    problems = []
    if meta is None:
        problems.append({"type": "format_error",
                         "detail": "no grading record in the file (not written by annotate.py, or "
                                   "edited afterwards) — its rubric and question tags are unknown"})
        return None, problems
    if meta.get("solutions_sha256") != cache.sha256_of(solutions):
        if meta.get("solutions_file") == cache.rel_path(solutions):
            problems.append({"type": "outdated_rubric",
                             "detail": "graded against an earlier version of this solutions file "
                                       "(it has changed since) — regrade"})
        else:
            problems.append({"type": "format_error",
                             "detail": "graded against a different solutions file than the one given"})
    try:
        current = rubric.verified_rubric(solutions)
    except rubric.RubricError as exc:
        problems.append({"type": "outdated_rubric", "detail": f"no usable rubric now: {exc}"})
    else:
        if meta.get("rubric_sha256") != current["rubric_sha256"]:
            problems.append({"type": "outdated_rubric",
                             "detail": f"graded against rubric round {meta.get('rubric_round')}, "
                                       f"which is no longer the verified one (now round "
                                       f"{current['round']}) — regrade"})
    return meta, problems


def load_blind(path: Path, question_ids: list[str]) -> dict:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"cannot read blind verdicts: {exc}")
    items = data.get("verdicts") if isinstance(data, dict) else None
    if not isinstance(items, list):
        raise ValueError('blind verdicts must be {"verdicts": [{"question", "verdict", …}]}')
    verdicts = {}
    for n, v in enumerate(items):
        if not isinstance(v, dict) or v.get("question") not in question_ids:
            raise ValueError(f"verdicts[{n}]: 'question' must be a rubric question id {question_ids}")
        if str(v.get("verdict", "")).lower() not in BLIND_VERDICTS:
            raise ValueError(f"verdicts[{n}]: 'verdict' must be one of {BLIND_VERDICTS}")
        verdicts[v["question"]] = str(v["verdict"]).lower()
    missing = [q for q in question_ids if q not in verdicts]
    if missing:
        raise ValueError(f"commit a blind verdict for every rubric question first — missing {missing}")
    return verdicts


def compare(blind: dict, annotations: list[dict], meta: dict | None) -> list[dict]:
    by_n = {a["n"]: a["question"] for a in (meta or {}).get("annotations", [])}
    grader: dict[str, set] = {}
    for a in annotations:
        question = by_n.get(a.get("tag"))
        a["question"] = question
        if question:
            grader.setdefault(question, set()).add((a["verdict_label"] or "?").lower())
    out = []
    for question, mine in blind.items():
        labels = grader.get(question, set())
        theirs = "correct" if not labels else "incorrect" if "incorrect" in labels else "incomplete"
        if mine == theirs:
            status = "agree"
        elif "correct" in (mine, theirs) or "unreadable" in (mine, theirs):
            status = "verdict_mismatch"
        else:
            status = "label_mismatch"
        out.append({"question": question, "checker": mine, "grader": theirs, "status": status})
    return out


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("graded_file", type=Path)
    parser.add_argument("submission_file", type=Path)
    parser.add_argument("--solutions", type=Path, required=True)
    parser.add_argument("--report", action="store_true",
                        help="produce the full report (needs --blind)")
    parser.add_argument("--blind", type=Path, help="the checker's committed blind verdicts")
    args = parser.parse_args()

    for p in (args.graded_file, args.submission_file, args.solutions):
        if not p.exists():
            print(json.dumps({"error": f"file not found: {p}"}))
            sys.exit(1)
    ext = args.graded_file.suffix.lower()
    if ext not in (".docx", ".xlsx"):
        print(json.dumps({"error": f"expected a _Graded.docx or _Graded.xlsx, got {ext}"}))
        sys.exit(1)
    try:
        state = marks.mark_state(args.graded_file)
        meta, prov = provenance_problems(args.graded_file, args.solutions)
        if not args.report:
            print(json.dumps({
                "grading_completed": state["grading_completed"],
                "review_mark_present": state["reviewed"],
                "rubric_problems": [p["detail"] for p in prov],
                "next": ("commit your blind verdicts, then rerun with --report --blind <file>"
                         if state["grading_completed"] else
                         "Grading not yet marked complete. Grader must add 'Grading Completed' "
                         "mark before checker can proceed."),
            }, indent=2, ensure_ascii=False))
            sys.exit(0 if state["grading_completed"] else 2)

        if args.blind is None:
            raise ValueError("--report needs --blind <checker_blind.json>: commit your own verdict "
                             "for every rubric question before seeing the grader's")
        # The checker grades against the current verified rubric (it cannot
        # grade blind without one); an outdated graded file shows up in prov.
        blind = load_blind(args.blind, rubric.verified_rubric(args.solutions)["question_ids"])
        report = (audit_document if ext == ".docx" else audit_workbook)(
            args.graded_file, args.submission_file)
        report["problems"] = prov + report["problems"]
        for a in report["annotations"]:
            if a.get("tag") is None:
                report["problems"].append({"type": "format_error", "annotation": a["n"],
                                           "detail": "annotation has no question tag"})
        report["comparison"] = compare(blind, report["annotations"], meta)
        report["blind_sha256"] = hashlib.sha256(args.blind.read_bytes()).hexdigest()
        report["graded_sha256"] = cache.sha256_of(args.graded_file)
    except (RuntimeError, ValueError, rubric.RubricError) as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False))
        sys.exit(1)
    except Exception as exc:  # corrupt/unreadable graded file, conversion failure, …
        print(json.dumps({"error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False))
        sys.exit(1)

    report_path = cache.work_dir_for(args.submission_file) / "audit.json"
    cache.write_text_atomic(report_path, json.dumps(report, indent=1, ensure_ascii=False))
    counts: dict[str, int] = {}
    for row in report["comparison"]:
        counts[row["status"]] = counts.get(row["status"], 0) + 1
    print(json.dumps({"report_path": str(report_path), "comparison": counts,
                      "mechanical_problems": len(report["problems"])}, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
