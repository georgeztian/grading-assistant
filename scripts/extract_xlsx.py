#!/usr/bin/env python3
"""Extract every tab of a .xlsx workbook — formulas, cached values, hidden
sheets/rows/columns, cell comments, and merged-cell ranges.

Implements the "Spreadsheet Extraction" method from grader.md: a workbook is
loaded twice (data_only=False for formulas, data_only=True for cached
values) since neither load alone gives both, and every sheet is walked
(including hidden ones) rather than assuming a single active tab.

Usage:
    .venv/Scripts/python scripts/extract_xlsx.py <file.xlsx>
    .venv/Scripts/python scripts/extract_xlsx.py <file.xlsx> --cache
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import _venv  # noqa: F401  — must precede third-party imports (re-runs under .venv)

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib import cache  # noqa: E402


def extract_xlsx(path: Path) -> dict:
    wb_formulas = openpyxl.load_workbook(str(path), data_only=False)
    wb_values = openpyxl.load_workbook(str(path), data_only=True)

    sheets = {}
    cells_with_formula = 0
    cells_with_comment = 0
    hidden_sheet_count = 0

    for sheet_name in wb_formulas.sheetnames:
        ws_f = wb_formulas[sheet_name]
        ws_v = wb_values[sheet_name]
        is_hidden = ws_f.sheet_state != "visible"
        if is_hidden:
            hidden_sheet_count += 1

        hidden_rows = [r for r, dim in ws_f.row_dimensions.items() if dim.hidden]
        hidden_cols = [c for c, dim in ws_f.column_dimensions.items() if dim.hidden]
        merged_ranges = [str(r) for r in ws_f.merged_cells.ranges]
        image_count = len(getattr(ws_f, "_images", []))  # only populated if Pillow is installed
        chart_count = len(getattr(ws_f, "_charts", []))

        cells: dict[str, dict] = {}
        max_row = max(ws_f.max_row, ws_v.max_row)
        max_col = max(ws_f.max_column, ws_v.max_column)
        for row_f, row_v in zip(
            ws_f.iter_rows(min_row=1, max_row=max_row, max_col=max_col),
            ws_v.iter_rows(min_row=1, max_row=max_row, max_col=max_col),
        ):
            for cell_f, cell_v in zip(row_f, row_v):
                raw = cell_f.value
                formula = raw if isinstance(raw, str) and raw.startswith("=") else None
                value = cell_v.value if formula is not None else raw
                comment = cell_f.comment.text if cell_f.comment else None

                if formula is None and value is None and comment is None:
                    continue

                if formula is not None:
                    cells_with_formula += 1
                if comment is not None:
                    cells_with_comment += 1

                entry = {}
                if value is not None:
                    entry["value"] = value if not hasattr(value, "isoformat") else value.isoformat()
                if formula is not None:
                    entry["formula"] = formula
                    if value is None:
                        entry["value_uncached"] = True
                if comment is not None:
                    entry["comment"] = comment
                cells[cell_f.coordinate] = entry

        sheets[sheet_name] = {
            "hidden": is_hidden,
            "hidden_rows": hidden_rows,
            "hidden_columns": hidden_cols,
            "merged_ranges": merged_ranges,
            "image_count": image_count,
            "chart_count": chart_count,
            "cells": cells,
        }

    total_images = sum(s["image_count"] for s in sheets.values())
    total_charts = sum(s["chart_count"] for s in sheets.values())

    return {
        "type": "xlsx",
        "sheet_order": wb_formulas.sheetnames,
        "sheets": sheets,
        "summary": {
            "sheet_count": len(wb_formulas.sheetnames),
            "hidden_sheet_count": hidden_sheet_count,
            "cells_with_formula": cells_with_formula,
            "cells_with_comment": cells_with_comment,
            "image_count": total_images,
            "chart_count": total_charts,
            "image_warning": (
                f"{total_images} embedded image(s) and {total_charts} chart(s) "
                "found across sheets (see per-sheet image_count/chart_count) — "
                "their content is NOT in this extraction (no OCR/chart-reading; "
                "a chart's source data is in cells, but its rendering, titles "
                "and labels are not). If a question depends on one, inspect it "
                f"directly, e.g.: `unzip -o \"{path}\" -d "
                f"\".cache/inspect/{path.stem}\"` then view "
                f"\".cache/inspect/{path.stem}/xl/media/\"* (images) or "
                f"\".cache/inspect/{path.stem}/xl/charts/\"* (charts) — use the "
                "project's existing .cache/ scratch area, not an ad hoc path."
            ) if (total_images or total_charts) else None,
        },
    }


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__,
                                      formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("file", type=Path)
    parser.add_argument("--cache", action="store_true")
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()

    if not args.file.exists():
        print(json.dumps({"error": f"file not found: {args.file}"}))
        sys.exit(1)

    if args.cache and not args.force:
        cached = cache.load_cached(args.file)
        if cached is not None:
            print(json.dumps({
                "cache_hit": True,
                "cache_path": str(cache.cache_path_for(args.file)),
                "summary": cached.get("summary"),
            }, indent=2))
            return

    record = extract_xlsx(args.file)

    if args.cache:
        cache_file = cache.save_cache(args.file, record)
        print(json.dumps({
            "cache_hit": False,
            "cache_path": str(cache_file),
            "summary": record["summary"],
        }, indent=2))
    else:
        print(json.dumps(record, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
