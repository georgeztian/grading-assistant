#!/usr/bin/env python3
"""Extract text + equations (in document order) from a .docx file.

Implements the "Equation Extraction" method from grader.md: python-docx's
`paragraph.text` silently returns "" for Word equation paragraphs (they are
stored as OMML XML, not `w:t` runs), so this walks each paragraph's XML
directly and inlines any `<m:oMath>` as a linearized "[EQ: ...]" expression
in the correct position relative to the surrounding text.

Usage:
    python extract_docx.py <file.docx>            # print JSON to stdout
    python extract_docx.py <file.docx> --cache     # write to the shared
                                                    # extraction cache and
                                                    # print only the path
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib import cache  # noqa: E402
from lib.omml import M_NS, linearize_omath  # noqa: E402

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def paragraph_text_with_equations(paragraph: Paragraph) -> tuple[str, list[str]]:
    """Render one paragraph's text in document order, inlining equations.

    Returns (text, equations) where `equations` lists each linearized
    equation found in the paragraph, in order.
    """
    parts: list[str] = []
    equations: list[str] = []

    def walk(el):
        tag = etree.QName(el).localname
        ns = etree.QName(el).namespace
        if ns == M_NS and tag == "oMath":
            expr = linearize_omath(el)
            equations.append(expr)
            parts.append(f"[EQ: {expr}]")
            return  # don't also descend into its w:t-less internals
        if ns == W_NS and tag == "t":
            parts.append(el.text or "")
            return
        if ns == W_NS and tag == "tab":
            parts.append("\t")
            return
        if ns == W_NS and tag == "br":
            parts.append("\n")
            return
        for child in el:
            walk(child)

    walk(paragraph._p)
    text = "".join(parts)
    return text, equations


def iter_block_items(parent):
    """Yield Paragraph/Table objects from a Document or _Cell in document
    order (python-docx has no built-in for this — paragraphs and tables are
    siblings in the underlying XML body)."""
    if hasattr(parent, "element"):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._tc
    for child in parent_elm.iterchildren():
        tag = etree.QName(child).localname
        if tag == "p":
            yield Paragraph(child, parent)
        elif tag == "tbl":
            yield Table(child, parent)


def extract_table(table: Table, table_index: int) -> dict:
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            cell_paragraphs = []
            for block in iter_block_items(cell):
                if isinstance(block, Paragraph):
                    text, eqs = paragraph_text_with_equations(block)
                    cell_paragraphs.append(text)
            cells.append("\n".join(cell_paragraphs))
        rows.append(cells)
    return {"type": "table", "docx_table_index": table_index, "rows": rows}


def count_images(document) -> int:
    """Count embedded images via the document's own part relationships —
    catches both inline and floating/anchored images, unlike
    `document.inline_shapes` (inline only). A nonzero count means the
    extracted text is NOT the full content: figures/tables baked into an
    image never appear as text or as a `table` block."""
    count = 0
    for rel in document.part.rels.values():
        if "image" in rel.reltype:
            count += 1
    return count


def extract_docx(path: Path) -> dict:
    document = docx.Document(str(path))
    blocks = []
    equations_found = 0
    paragraph_index = 0  # tracks document.paragraphs[i] — top-level paragraphs only
    table_index = 0  # tracks document.tables[i]
    for block in iter_block_items(document):
        if isinstance(block, Paragraph):
            text, eqs = paragraph_text_with_equations(block)
            equations_found += len(eqs)
            blocks.append({
                "type": "paragraph",
                "docx_paragraph_index": paragraph_index,
                "text": text,
                "style": block.style.name if block.style else None,
                "is_empty": not text.strip(),
                "has_equation": bool(eqs),
                "equations": eqs,
            })
            paragraph_index += 1
        elif isinstance(block, Table):
            blocks.append(extract_table(block, table_index))
            table_index += 1

    image_count = count_images(document)

    return {
        "type": "docx",
        "blocks": blocks,
        "summary": {
            "paragraph_count": sum(1 for b in blocks if b["type"] == "paragraph"),
            "table_count": sum(1 for b in blocks if b["type"] == "table"),
            "equations_found": equations_found,
            "non_empty_paragraphs": sum(
                1 for b in blocks if b["type"] == "paragraph" and not b["is_empty"]
            ),
            "image_count": image_count,
            "image_warning": (
                f"{image_count} embedded image(s) found — their content is NOT "
                "in this extraction (no OCR). If a question, figure, or table "
                "depends on an image, inspect it directly, e.g.: "
                "`unzip -o file.docx -d /tmp/x` then view /tmp/x/word/media/*."
            ) if image_count else None,
        },
        "notes": {
            "docx_paragraph_index": "Index into python-docx's `document.paragraphs` "
                "(top-level body paragraphs only, tables excluded) — use this to "
                "locate the exact Paragraph object for inserting an annotation "
                "after it, instead of re-matching on text.",
            "docx_table_index": "Index into `document.tables`, in document order.",
        },
    }


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__,
                                      formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("file", type=Path)
    parser.add_argument("--cache", action="store_true",
                         help="write result to the shared extraction cache "
                              "and print only the cache file path + summary")
    parser.add_argument("--force", action="store_true",
                         help="ignore any existing cache entry and re-extract")
    args = parser.parse_args()

    if not args.file.exists():
        print(json.dumps({"error": f"file not found: {args.file}"}))
        sys.exit(1)

    if args.cache and not args.force:
        cached = cache.load_cached(args.file)
        if cached is not None:
            print(json.dumps({
                "cache_hit": True,
                "cache_path": str(cache.cache_path_for(args.file)),
                "summary": cached.get("summary"),
            }, indent=2))
            return

    record = extract_docx(args.file)

    if args.cache:
        cache_file = cache.save_cache(args.file, record)
        print(json.dumps({
            "cache_hit": False,
            "cache_path": str(cache_file),
            "summary": record["summary"],
        }, indent=2))
    else:
        print(json.dumps(record, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
