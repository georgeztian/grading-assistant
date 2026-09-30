"""Evaluate one Excel formula against chosen cell values — enough of Excel to
re-run a student's formula on the solution's inputs, which tells whether the
formula itself is right independently of any upstream error.

`evaluate(formula, sheet, lookup)`: `lookup(sheet, "C12")` returns a cell's
value (None when empty). Anything outside the supported subset (named
ranges, whole-column references, array constants, unknown functions, …)
raises Unsupported; an Excel error result (#DIV/0!, #VALUE!, #NUM!, …) raises
ExcelError. Callers treat both as "cannot decide by evaluation".
"""
from __future__ import annotations

import math
import re
from decimal import ROUND_DOWN, ROUND_HALF_UP, ROUND_UP, Decimal, InvalidOperation

from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.utils.cell import column_index_from_string, get_column_letter

MAX_RANGE_CELLS = 10000
REF_RE = re.compile(
    r"^(?:(?P<sheet>'(?:[^']|'')+'|[^'!:]+)!)?"
    r"\$?(?P<c1>[A-Za-z]{1,3})\$?(?P<r1>[0-9]+)"
    r"(?::\$?(?P<c2>[A-Za-z]{1,3})\$?(?P<r2>[0-9]+))?$")
PRECEDENCE = {"=": 1, "<>": 1, "<": 1, ">": 1, "<=": 1, ">=": 1,
              "&": 2, "+": 3, "-": 3, "*": 4, "/": 4, "^": 5}
ERROR_VALUES = {"#DIV/0!", "#VALUE!", "#NUM!", "#N/A", "#REF!", "#NAME?", "#NULL!"}


class Unsupported(Exception):
    pass


class ExcelError(Exception):
    pass


class Range:
    """A multi-cell reference, as a function argument sees it."""

    def __init__(self, rows: list[list]):
        self.rows = rows

    def flat(self) -> list:
        return [v for row in self.rows for v in row]


# ------------------------------------------------------------------ parsing

def parse_ref(text: str, sheet: str) -> tuple[str, int, int, int, int]:
    m = REF_RE.match(text)
    if not m:
        raise Unsupported(f"reference {text!r} (named range, whole column/row or 3-D reference)")
    ref_sheet = m.group("sheet")
    if ref_sheet is None:
        ref_sheet = sheet
    elif ref_sheet.startswith("'"):
        ref_sheet = ref_sheet[1:-1].replace("''", "'")
    if ref_sheet is not None and ("[" in ref_sheet or "]" in ref_sheet):
        # Excel forbids [ ] in sheet names: this is another workbook ([1]Sheet1!A1)
        raise Unsupported(f"reference {text!r} to another workbook")
    c1, r1 = column_index_from_string(m.group("c1").upper()), int(m.group("r1"))
    c2 = column_index_from_string(m.group("c2").upper()) if m.group("c2") else c1
    r2 = int(m.group("r2")) if m.group("r2") else r1
    return ref_sheet, min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2)


def references(formula: str, sheet: str) -> list[tuple[str, str]]:
    """Every (sheet, cell) a formula reads, ranges expanded; unparseable
    references are skipped."""
    out = []
    try:
        items = Tokenizer(formula).items
    except Exception:
        return out
    for tok in items:
        if tok.type == Token.OPERAND and tok.subtype == Token.RANGE:
            try:
                ref_sheet, r1, c1, r2, c2 = parse_ref(tok.value, sheet)
            except Unsupported:
                continue
            if (r2 - r1 + 1) * (c2 - c1 + 1) > MAX_RANGE_CELLS:
                continue
            out += [(ref_sheet, f"{get_column_letter(c)}{r}")
                    for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
    return out


class Parser:
    def __init__(self, formula: str):
        if not isinstance(formula, str) or not formula.startswith("="):
            raise Unsupported("not a formula")
        try:
            items = Tokenizer(formula).items
        except Exception as exc:
            raise Unsupported(f"cannot tokenize: {exc}")
        self.toks = [t for t in items if t.type != Token.WSPACE]
        self.i = 0

    def peek(self):
        return self.toks[self.i] if self.i < len(self.toks) else None

    def take(self):
        tok = self.peek()
        if tok is None:
            raise Unsupported("unexpected end of formula")
        self.i += 1
        return tok

    def parse(self):
        if not self.toks:
            raise Unsupported("empty formula")
        node = self.expr(1)
        if self.peek() is not None:
            raise Unsupported(f"unexpected {self.peek().value!r}")
        return node

    def expr(self, min_prec: int):
        left = self.unary()
        while True:
            tok = self.peek()
            if tok is None or tok.type != Token.OP_IN or PRECEDENCE.get(tok.value, 0) < min_prec:
                return left
            if tok.value not in PRECEDENCE:
                raise Unsupported(f"operator {tok.value!r}")
            self.take()
            left = ("bin", tok.value, left, self.expr(PRECEDENCE[tok.value] + 1))

    def unary(self):
        tok = self.peek()
        if tok is not None and tok.type == Token.OP_PRE:
            self.take()
            operand = self.unary()
            return ("neg", operand) if tok.value == "-" else ("pos", operand)
        node = self.primary()
        while self.peek() is not None and self.peek().type == Token.OP_POST:
            self.take()
            node = ("pct", node)
        return node

    def primary(self):
        tok = self.take()
        if tok.type == Token.OPERAND:
            if tok.subtype == Token.NUMBER:
                return ("lit", float(tok.value))
            if tok.subtype == Token.TEXT:
                return ("lit", tok.value[1:-1].replace('""', '"'))
            if tok.subtype == Token.LOGICAL:
                return ("lit", tok.value.upper() == "TRUE")
            if tok.subtype == Token.ERROR:
                return ("err", tok.value)
            return ("ref", tok.value)
        if tok.type == Token.PAREN and tok.subtype == Token.OPEN:
            node = self.expr(1)
            close = self.take()
            if close.type != Token.PAREN or close.subtype != Token.CLOSE:
                raise Unsupported("unbalanced parentheses")
            return node
        if tok.type == Token.FUNC and tok.subtype == Token.OPEN:
            name = tok.value[:-1].upper()
            for prefix in ("_XLFN.", "_XLWS."):
                if name.startswith(prefix):
                    name = name[len(prefix):]
            args = []
            nxt = self.peek()
            if nxt is not None and nxt.type == Token.FUNC and nxt.subtype == Token.CLOSE:
                self.take()
                return ("fn", name, args)
            while True:
                nxt = self.peek()
                if nxt is not None and (nxt.type == Token.SEP and nxt.subtype == Token.ARG
                                        or nxt.type == Token.FUNC and nxt.subtype == Token.CLOSE):
                    args.append(None)  # omitted argument, e.g. PV(r,n,,fv)
                else:
                    args.append(self.expr(1))
                sep = self.take()
                if sep.type == Token.FUNC and sep.subtype == Token.CLOSE:
                    return ("fn", name, args)
                if not (sep.type == Token.SEP and sep.subtype == Token.ARG):
                    raise Unsupported(f"unexpected {sep.value!r} in {name}()")
        raise Unsupported(f"unsupported syntax {tok.value!r}")


# ------------------------------------------------------------------ values

def is_error(value) -> bool:
    return isinstance(value, str) and value in ERROR_VALUES


def to_number(value) -> float:
    if is_error(value):
        raise ExcelError(value)
    if value is None:
        return 0.0
    if isinstance(value, bool):
        return 1.0 if value else 0.0
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.strip().replace(",", ""))
        except ValueError:
            raise ExcelError("#VALUE!")
    raise ExcelError("#VALUE!")


def to_text(value) -> str:
    if is_error(value):
        raise ExcelError(value)
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return str(int(value)) if float(value).is_integer() else format(value, ".15g")
    return str(value)


def to_bool(value) -> bool:
    if isinstance(value, str) and not is_error(value):
        if value.upper() in ("TRUE", "FALSE"):
            return value.upper() == "TRUE"
        raise ExcelError("#VALUE!")
    return to_number(value) != 0


def compare(op: str, a, b) -> bool:
    def rank(v):
        if isinstance(v, bool):
            return 2
        if isinstance(v, str):
            return 1
        return 0
    if a is None:
        a = "" if isinstance(b, str) else False if isinstance(b, bool) else 0.0
    if b is None:
        b = "" if isinstance(a, str) else False if isinstance(a, bool) else 0.0
    for v in (a, b):
        if is_error(v):
            raise ExcelError(v)
    if rank(a) != rank(b):
        a, b = rank(a), rank(b)
    elif isinstance(a, str):
        a, b = a.lower(), b.lower()
    elif rank(a) == 0:
        # Excel compares numbers at 15 significant digits (=0.1+0.2=0.3 is TRUE)
        a, b = float(format(a, ".15g")), float(format(b, ".15g"))
    return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b,
            "<=": a <= b, ">=": a >= b}[op]


def excel_round(x: float, digits: float, mode=ROUND_HALF_UP) -> float:
    try:
        return float(Decimal(repr(float(x))).quantize(Decimal(1).scaleb(-int(digits)), rounding=mode))
    except InvalidOperation:
        return float(x)


# ------------------------------------------------------------------ evaluation

class Evaluator:
    def __init__(self, sheet: str, lookup):
        self.sheet = sheet
        self.lookup = lookup

    def ref(self, text: str):
        ref_sheet, r1, c1, r2, c2 = parse_ref(text, self.sheet)
        if r1 == r2 and c1 == c2:
            return self.lookup(ref_sheet, f"{get_column_letter(c1)}{r1}")
        if (r2 - r1 + 1) * (c2 - c1 + 1) > MAX_RANGE_CELLS:
            raise Unsupported(f"range {text} is too large")
        return Range([[self.lookup(ref_sheet, f"{get_column_letter(c)}{r}")
                       for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)])

    def value(self, node):
        """A node's value as a scalar (a multi-cell range is not a scalar)."""
        v = self.raw(node)
        if isinstance(v, Range):
            raise Unsupported("a multi-cell range used as a single value")
        return v

    def raw(self, node):
        kind = node[0]
        if kind == "lit":
            return node[1]
        if kind == "err":
            raise ExcelError(node[1])
        if kind == "ref":
            return self.ref(node[1])
        if kind == "neg":
            return -to_number(self.value(node[1]))
        if kind == "pos":
            return self.value(node[1])
        if kind == "pct":
            return to_number(self.value(node[1])) / 100
        if kind == "bin":
            return self.binary(node[1], self.value(node[2]), self.value(node[3]))
        if kind == "fn":
            return self.call(node[1], node[2])
        raise Unsupported(kind)

    @staticmethod
    def binary(op: str, a, b):
        if op == "&":
            return to_text(a) + to_text(b)
        if op in ("=", "<>", "<", ">", "<=", ">="):
            return compare(op, a, b)
        x, y = to_number(a), to_number(b)
        if op == "+":
            return x + y
        if op == "-":
            return x - y
        if op == "*":
            return x * y
        if op == "/":
            if y == 0:
                raise ExcelError("#DIV/0!")
            return x / y
        if op == "^":
            if x == 0 and y <= 0:
                raise ExcelError("#NUM!" if y == 0 else "#DIV/0!")
            if x < 0 and not float(y).is_integer():
                raise ExcelError("#NUM!")
            try:
                return x ** y
            except OverflowError:
                raise ExcelError("#NUM!")
        raise Unsupported(f"operator {op}")

    # -- argument helpers
    def numbers(self, args, skip_errors: bool = False) -> list[float]:
        """SUM-style: typed arguments are coerced; referenced cells count only
        when they hold numbers (text, logicals and blanks are skipped).
        `skip_errors` (COUNT): errors, and typed arguments that are not
        numbers, are ignored instead of raised."""
        out = []
        for a in args:
            if a is None:
                continue
            try:
                v = self.raw(a)
            except ExcelError:
                if skip_errors:
                    continue
                raise
            if a[0] == "ref" and not isinstance(v, Range):
                v = Range([[v]])
            if isinstance(v, Range):
                for x in v.flat():
                    if is_error(x):
                        if skip_errors:
                            continue
                        raise ExcelError(x)
                    if isinstance(x, (int, float)) and not isinstance(x, bool):
                        out.append(float(x))
            else:
                try:
                    out.append(to_number(v))
                except ExcelError:
                    if not skip_errors:
                        raise
        return out

    def arg(self, args, n: int, default=None):
        if n < len(args) and args[n] is not None:
            return self.value(args[n])
        if default is None:
            raise ExcelError("#VALUE!")
        return default

    def num(self, args, n: int, default=None) -> float:
        return to_number(self.arg(args, n, default))

    def pairs(self, args) -> tuple[list[float], list[float]]:
        if len(args) != 2:
            raise Unsupported("expected two ranges")
        a, b = (self.raw(x) for x in args)
        a = a.flat() if isinstance(a, Range) else [a]
        b = b.flat() if isinstance(b, Range) else [b]
        if len(a) != len(b):
            raise ExcelError("#N/A")
        xs, ys = [], []
        for x, y in zip(a, b):
            if isinstance(x, (int, float)) and isinstance(y, (int, float)) \
                    and not isinstance(x, bool) and not isinstance(y, bool):
                xs.append(float(x))
                ys.append(float(y))
        return xs, ys

    def call(self, name: str, args: list):
        fn = FUNCTIONS.get(name)
        if fn is None:
            raise Unsupported(f"function {name}")
        return fn(self, args)


def _mean(xs):
    if not xs:
        raise ExcelError("#DIV/0!")
    return sum(xs) / len(xs)


def _var(xs, sample: bool):
    n = len(xs)
    if n < (2 if sample else 1):
        raise ExcelError("#DIV/0!")
    m = sum(xs) / n
    return sum((x - m) ** 2 for x in xs) / (n - 1 if sample else n)


def _cov(xs, ys, sample: bool):
    n = len(xs)
    if n < (2 if sample else 1):
        raise ExcelError("#DIV/0!")
    mx, my = sum(xs) / n, sum(ys) / n
    return sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / (n - 1 if sample else n)


def _slope(ev, args):
    ys, xs = ev.pairs(args)  # SLOPE(known_y, known_x)
    vx = _var(xs, False)
    if vx == 0:
        raise ExcelError("#DIV/0!")
    return _cov(xs, ys, False) / vx


def _intercept(ev, args):
    ys, xs = ev.pairs(args)
    return _mean(ys) - _slope(ev, args) * _mean(xs)


def _correl(ev, args):
    xs, ys = ev.pairs(args)
    d = math.sqrt(_var(xs, False) * _var(ys, False))
    if d == 0:
        raise ExcelError("#DIV/0!")
    return _cov(xs, ys, False) / d


def _pv(rate, nper, pmt, fv=0.0, typ=0.0):
    if rate == 0:
        return -(fv + pmt * nper)
    g = (1 + rate) ** nper
    return -(fv + pmt * (1 + rate * typ) * (g - 1) / rate) / g


def _fv(rate, nper, pmt, pv=0.0, typ=0.0):
    if rate == 0:
        return -(pv + pmt * nper)
    g = (1 + rate) ** nper
    return -(pv * g + pmt * (1 + rate * typ) * (g - 1) / rate)


def _pmt(rate, nper, pv, fv=0.0, typ=0.0):
    if nper == 0:
        raise ExcelError("#NUM!")
    if rate == 0:
        return -(pv + fv) / nper
    g = (1 + rate) ** nper
    return -(rate * (fv + pv * g)) / ((1 + rate * typ) * (g - 1))


def _nper(rate, pmt, pv, fv=0.0, typ=0.0):
    if rate == 0:
        if pmt == 0:
            raise ExcelError("#NUM!")
        return -(pv + fv) / pmt
    num = pmt * (1 + rate * typ) - fv * rate
    den = pmt * (1 + rate * typ) + pv * rate
    if den == 0 or num / den <= 0:
        raise ExcelError("#NUM!")
    return math.log(num / den) / math.log(1 + rate)


def _solve(f, guess: float) -> float:
    """A root of f near `guess` (Newton, then a bracketing scan)."""
    r = guess
    for _ in range(100):
        try:
            y = f(r)
            h = 1e-7 * max(1.0, abs(r))
            d = (f(r + h) - y) / h
        except (ZeroDivisionError, OverflowError, ValueError):
            break
        if d == 0:
            break
        step = y / d
        r -= step
        if r <= -1:
            break
        if abs(step) < 1e-12 * max(1.0, abs(r)):
            return r
    grid = [-0.99 + k * 0.01 for k in range(0, 200)] + [1 + k * 0.5 for k in range(1, 40)]
    prev = None
    for x in grid:
        try:
            y = f(x)
        except (ZeroDivisionError, OverflowError, ValueError):
            prev = None
            continue
        if prev is not None and (prev[1] < 0) != (y < 0):
            lo, hi = prev[0], x
            for _ in range(200):
                mid = (lo + hi) / 2
                if (f(lo) < 0) != (f(mid) < 0):
                    hi = mid
                else:
                    lo = mid
            return (lo + hi) / 2
        prev = (x, y)
    raise ExcelError("#NUM!")


def _fin_args(ev, args, required: int, optional: int) -> list[float]:
    """Financial-function arguments: the first `required` must be present
    (an omitted one, as in PV(r,n,,fv), counts as 0, like Excel)."""
    if len(args) < required or len(args) > required + optional:
        raise ExcelError("#VALUE!")
    vals = [ev.num(args, k, 0.0) for k in range(len(args))]
    if len(vals) > required + 1:  # `type`: any nonzero value means "beginning of period"
        vals[required + 1] = 1.0 if vals[required + 1] else 0.0
    return vals


def _rate(ev, args):
    vals = _fin_args(ev, args, 3, 3)
    nper, pmt, pv = vals[:3]
    fv = vals[3] if len(vals) > 3 else 0.0
    typ = vals[4] if len(vals) > 4 else 0.0
    guess = vals[5] if len(vals) > 5 and args[5] is not None else 0.1

    def f(r):
        if r == 0:
            return pv + pmt * nper + fv
        g = (1 + r) ** nper
        return pv * g + pmt * (1 + r * typ) * (g - 1) / r + fv
    return _solve(f, guess)


def _npv(ev, args):
    rate = ev.num(args, 0)
    values = ev.numbers(args[1:])
    return sum(v / (1 + rate) ** (i + 1) for i, v in enumerate(values))


def _irr(ev, args):
    values = ev.numbers(args[:1])
    guess = ev.num(args, 1, 0.1)
    if not any(v > 0 for v in values) or not any(v < 0 for v in values):
        raise ExcelError("#NUM!")
    return _solve(lambda r: sum(v / (1 + r) ** i for i, v in enumerate(values)), guess)


def _if(ev, args):
    if not args or args[0] is None:
        raise ExcelError("#VALUE!")
    if to_bool(ev.value(args[0])):
        return ev.value(args[1]) if len(args) > 1 and args[1] is not None else 0.0
    if len(args) > 2:
        return ev.value(args[2]) if args[2] is not None else 0.0
    return False


def _iferror(ev, args):
    if not args or args[0] is None:
        raise ExcelError("#VALUE!")
    try:
        v = ev.value(args[0])
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            raise ExcelError("#NUM!")
        return v
    except (ExcelError, ZeroDivisionError, OverflowError, ValueError):  # math errors are #NUM!
        return ev.value(args[1]) if len(args) > 1 and args[1] is not None else 0.0


def _logical(combine):
    def fn(ev, args):
        vals = []
        for a in args:
            if a is None:
                continue
            v = ev.raw(a)
            if a[0] == "ref" and not isinstance(v, Range):
                v = Range([[v]])  # a referenced cell's text or blank is ignored
            if isinstance(v, Range):
                for x in v.flat():
                    if is_error(x):
                        raise ExcelError(x)
                vals += [bool(x) for x in v.flat() if isinstance(x, (bool, int, float))]
            else:
                vals.append(to_bool(v))
        if not vals:
            raise ExcelError("#VALUE!")
        return combine(vals)
    return fn


def _sumproduct(ev, args):
    arrays = []
    for a in args:
        if a is None:
            raise ExcelError("#VALUE!")
        v = ev.raw(a)
        arrays.append(v.flat() if isinstance(v, Range) else [v])
    if len({len(x) for x in arrays}) != 1:
        raise ExcelError("#VALUE!")
    total = 0.0
    for items in zip(*arrays):
        prod = 1.0
        for x in items:
            if is_error(x):
                raise ExcelError(x)
            prod *= float(x) if isinstance(x, (int, float)) and not isinstance(x, bool) else 0.0
        total += prod
    return total


def _product(xs):
    return math.prod(xs) if xs else 0.0  # Excel: PRODUCT of no numbers is 0


def _median(xs):
    if not xs:
        raise ExcelError("#NUM!")
    s, n = sorted(xs), len(xs)
    return s[n // 2] if n % 2 else (s[n // 2 - 1] + s[n // 2]) / 2


def _log(ev, args):
    x, base = ev.num(args, 0), ev.num(args, 1, 10.0)
    if x <= 0 or base <= 0 or base == 1:
        raise ExcelError("#NUM!")
    return math.log(x, base)


def _positive(fn):
    def wrapped(x):
        if x <= 0:
            raise ExcelError("#NUM!")
        return fn(x)
    return wrapped


def _mod(a, b):
    if b == 0:
        raise ExcelError("#DIV/0!")
    return a - b * math.floor(a / b)


def _sqrt(x):
    if x < 0:
        raise ExcelError("#NUM!")
    return math.sqrt(x)


def _one(fn):
    return lambda ev, args: fn(ev.num(args, 0))


def _two(fn):
    return lambda ev, args: fn(ev.num(args, 0), ev.num(args, 1))


def _fin(fn, required: int):
    return lambda ev, args: fn(*_fin_args(ev, args, required, 2))


FUNCTIONS = {
    "SUM": lambda ev, args: sum(ev.numbers(args)),
    "AVERAGE": lambda ev, args: _mean(ev.numbers(args)),
    "MIN": lambda ev, args: min(ev.numbers(args), default=0.0),
    "MAX": lambda ev, args: max(ev.numbers(args), default=0.0),
    "COUNT": lambda ev, args: float(len(ev.numbers(args, skip_errors=True))),
    "PRODUCT": lambda ev, args: _product(ev.numbers(args)),
    "MEDIAN": lambda ev, args: _median(ev.numbers(args)),
    "STDEV": lambda ev, args: math.sqrt(_var(ev.numbers(args), True)),
    "STDEV.S": lambda ev, args: math.sqrt(_var(ev.numbers(args), True)),
    "STDEV.P": lambda ev, args: math.sqrt(_var(ev.numbers(args), False)),
    "STDEVP": lambda ev, args: math.sqrt(_var(ev.numbers(args), False)),
    "VAR": lambda ev, args: _var(ev.numbers(args), True),
    "VAR.S": lambda ev, args: _var(ev.numbers(args), True),
    "VAR.P": lambda ev, args: _var(ev.numbers(args), False),
    "VARP": lambda ev, args: _var(ev.numbers(args), False),
    "COVAR": lambda ev, args: _cov(*ev.pairs(args), False),
    "COVARIANCE.P": lambda ev, args: _cov(*ev.pairs(args), False),
    "COVARIANCE.S": lambda ev, args: _cov(*ev.pairs(args), True),
    "CORREL": _correl,
    "SLOPE": _slope,
    "INTERCEPT": _intercept,
    "SUMPRODUCT": _sumproduct,
    "ABS": _one(abs),
    "SQRT": _one(_sqrt),
    "EXP": _one(math.exp),
    "LN": _one(_positive(math.log)),
    "LOG10": _one(_positive(math.log10)),
    "LOG": _log,
    "INT": _one(lambda x: float(math.floor(x))),
    "SIGN": _one(lambda x: float((x > 0) - (x < 0))),
    "POWER": lambda ev, args: Evaluator.binary("^", ev.num(args, 0), ev.num(args, 1)),
    "MOD": _two(_mod),
    "ROUND": _two(excel_round),
    "ROUNDUP": _two(lambda x, d: excel_round(x, d, ROUND_UP)),
    "ROUNDDOWN": _two(lambda x, d: excel_round(x, d, ROUND_DOWN)),
    "IF": _if,
    "IFERROR": _iferror,
    "AND": _logical(all),
    "OR": _logical(any),
    "NOT": lambda ev, args: not to_bool(ev.arg(args, 0)),
    "TRUE": lambda ev, args: True,
    "FALSE": lambda ev, args: False,
    "PV": _fin(_pv, 3),
    "FV": _fin(_fv, 3),
    "PMT": _fin(_pmt, 3),
    "NPER": _fin(_nper, 3),
    "RATE": _rate,
    "NPV": _npv,
    "IRR": _irr,
}


def evaluate(formula: str, sheet: str, lookup):
    """The formula's value, reading cells through `lookup(sheet, coord)`."""
    node = Parser(formula).parse()
    ev = Evaluator(sheet, lookup)
    try:
        result = ev.value(node)
    except (ZeroDivisionError, OverflowError, ValueError):  # math domain errors
        raise ExcelError("#NUM!")
    except (Unsupported, ExcelError):
        raise
    except Exception as exc:  # a construct the evaluator doesn't model (e.g. an omitted argument)
        raise Unsupported(f"{type(exc).__name__}: {exc}")
    if isinstance(result, float) and (math.isnan(result) or math.isinf(result)):
        raise ExcelError("#NUM!")
    return result


def normalize(formula: str | None) -> str | None:
    """A formula's text for an exact-equivalence check: no leading '=', no
    '$', no whitespace or case differences outside string literals."""
    if not formula:
        return None
    parts = re.split(r'("(?:[^"]|"")*")', formula.lstrip("="))
    out = []
    for n, part in enumerate(parts):
        if n % 2:
            out.append(part)
        else:
            part = re.sub(r"\s+", "", part.replace("$", "")).upper()
            out.append(part.replace("_XLFN.", "").replace("_XLWS.", ""))
    return "".join(out)
