#!/usr/bin/env python3
"""Write the checker's single-pass verdict into a `_Graded.docx|xlsx` — the
one command a checker runs instead of hand-writing python-docx/openpyxl code.

Everything grading-checker.md pins exactly is done here: blue (0000FF) text,
the exact wording `Review Passed` / `Review FAILED` (no variants, no
decoration), placement (after the red `Grading Completed` paragraph, or the
next row of the "Grading Summary" tab), and — for a failed review — one
blue line per discrepancy directly below it plus a closing summary line.

It will only mark a file the checker has audited blind: it requires the
report audit_graded.py wrote from the checker's committed blind verdicts
(`--report --blind`), for this graded file as it is now. `passed` is refused
while that report shows any verdict disagreement or mechanical problem, and a
`failed` review must list (at least) every disagreement the report found and
every kind (`type`) of mechanical problem in it.

discrepancies.json (write it in the submission's work_dir), for `failed`:
    {"discrepancies": [
        {"question": "Q2", "type": "verdict_mismatch",
         "checker_verdict": "incorrect", "checker_explanation": "...",
         "grader_verdict": "correct", "grader_explanation": "(no annotation)",
         "notes": "..."}
     ],
     "summary": "Grader's verdicts agree with the checker on every question except Q2 ..."}
  For a workbook graded cell by cell, a verdict disagreement names its
  `cell` ("Ex TN5 WACC!G12") instead of a question.
  type: one of lib/marks.py DISCREPANCY_TYPES (verdict_mismatch, label_mismatch,
  outdated_rubric, missed_question, annotation_placement, explanation_error,
  incomplete_coverage, cell_highlight_missing, missing_keypoint_not_flagged,
  format_error).
  question (or cell) / checker_verdict / grader_verdict are required for
  verdict_mismatch, label_mismatch and missed_question; a file-level problem
  (e.g. outdated_rubric) may leave them out but must describe itself in notes.

Usage:
    .venv/Scripts/python scripts/mark_review.py <graded_file> passed --submission <submission>
    .venv/Scripts/python scripts/mark_review.py <graded_file> failed <discrepancies.json> --submission <submission>
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

import _venv  # noqa: F401  — must precede third-party imports (re-runs under .venv)

import docx
import openpyxl
from docx.shared import RGBColor
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib import cache, marks  # noqa: E402
from lib.xlsx_shapes import restore_shapes  # noqa: E402
import annotate  # noqa: E402

FIELDS = ("question", "cell", "type", "checker_verdict", "checker_explanation",
          "grader_verdict", "grader_explanation", "notes")
VERDICT_TYPES = ("verdict_mismatch", "label_mismatch", "missed_question")


class ReviewError(Exception):
    pass


def load_discrepancies(path: Path) -> tuple[list[dict], str]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ReviewError(f"cannot read discrepancies JSON: {exc}")
    items = data.get("discrepancies") if isinstance(data, dict) else None
    if not isinstance(items, list) or not items:
        raise ReviewError("a failed review needs at least one entry in 'discrepancies'")
    for n, d in enumerate(items):
        if not isinstance(d, dict) or d.get("type") not in marks.DISCREPANCY_TYPES:
            raise ReviewError(f"discrepancies[{n}]: 'type' must be one of {list(marks.DISCREPANCY_TYPES)}")
        # A verdict disagreement names both verdicts; a file-level problem
        # (outdated answer key, altered content, …) may have no question or
        # verdicts, but must say what is wrong.
        required = (("checker_verdict", "grader_verdict") if d["type"] in VERDICT_TYPES else ())
        if d["type"] in VERDICT_TYPES and not (str(d.get("question", "")).strip()
                                               or str(d.get("cell", "")).strip()):
            raise ReviewError(f"discrepancies[{n}]: 'question' (or, for a workbook cell, 'cell') "
                              f"is required for {d['type']}")
        for field in required:
            if not str(d.get(field, "")).strip():
                raise ReviewError(f"discrepancies[{n}]: '{field}' is required for {d['type']}")
        if not any(str(d.get(f, "")).strip() for f in ("checker_verdict", "checker_explanation", "notes")):
            raise ReviewError(f"discrepancies[{n}]: state the problem in 'notes' (or the checker fields)")
    summary = str(data.get("summary", "")).strip()
    if not summary:
        raise ReviewError("'summary' (one closing line on agreement/disagreement — never a score) "
                          "is required")
    return items, summary


def discrepancy_line(d: dict) -> str:
    line = f"{d.get('cell') or d.get('question') or '(whole file)'} — {d['type']}:"
    for who in ("checker", "grader"):
        verdict, why = d.get(f"{who}_verdict"), d.get(f"{who}_explanation")
        if verdict or why:
            line += f" {who} says {verdict or ''}" + (f" ({why})" if why else "") + ";"
    line = line.rstrip(";")
    if d.get("notes"):
        line += f" {d['notes']}"
    return line if line.endswith(".") else line + "."


def check_audit(graded: Path, submission: Path, passed: bool, items: list[dict]) -> None:
    """The review must rest on a blind audit of this exact file."""
    report_path = cache.work_dir_for(submission) / "audit.json"
    if not report_path.exists():
        raise ReviewError("no audit report — run audit_graded.py --report --blind <your blind "
                          "verdicts> first")
    try:
        report = json.loads(report_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ReviewError(f"cannot read the audit report ({exc}) — run audit_graded.py --report again")
    if not report.get("blind_sha256"):
        raise ReviewError("the audit report was not built from committed blind verdicts")
    if report.get("graded_sha256") != cache.sha256_of(graded):
        raise ReviewError("the graded file changed after the audit — audit it again")
    disagreements = [c for c in report.get("comparison", []) if c["status"] != "agree"]
    if passed:
        if disagreements or report.get("problems"):
            raise ReviewError(f"cannot pass: the audit shows {len(disagreements)} verdict "
                              f"disagreement(s) and {len(report.get('problems', []))} problem(s) — "
                              "review them and mark 'failed' with each one listed")
        return
    listed = ({(d.get("question"), d.get("type")) for d in items}
              | {(annotate.cell_id(d["cell"]), d.get("type")) for d in items
                 if "!" in str(d.get("cell", ""))})
    # A comparison row is keyed by question (document) or cell (workbook). A
    # verdict_mismatch where the grader skipped the question may be reported
    # as missed_question instead.
    unlisted = [f"{c.get('question') or c.get('cell')} ({c['status']})" for c in disagreements
                if (c.get("question") or c.get("cell"), c["status"]) not in listed
                and not (c["status"] == "verdict_mismatch"
                         and (c.get("question") or c.get("cell"), "missed_question") in listed)]
    # Every mechanical problem is a real discrepancy: each kind the audit
    # found must be reported at least once.
    listed_types = {d.get("type") for d in items}
    unlisted += sorted({f"a '{p['type']}' problem ({p.get('detail', '')[:80]})"
                        for p in report.get("problems", []) if p.get("type") not in listed_types})
    if unlisted:
        raise ReviewError("these audit findings are not in discrepancies.json: "
                          + "; ".join(unlisted))


def check_state(state: dict) -> None:
    if not state["grading_completed"]:
        raise ReviewError("Grading not yet marked complete. Grader must add 'Grading Completed' "
                          "mark before checker can proceed.")
    if state["reviewed"]:
        raise ReviewError("file already carries a review mark — single pass only")


def mark_document(path: Path, passed: bool, items: list[dict], summary: str) -> None:
    document = docx.Document(str(path))
    blue = RGBColor.from_string(marks.CHECKER_BLUE)

    def add(text: str, bold: bool = False) -> None:
        run = document.add_paragraph().add_run(text)
        run.bold = bold
        run.font.color.rgb = blue

    add(marks.REVIEW_PASSED if passed else marks.REVIEW_FAILED, bold=True)
    for d in items:
        add(discrepancy_line(d))
    if not passed:
        add(summary)
    document.save(str(path))


def mark_workbook(path: Path, passed: bool, items: list[dict], summary: str) -> None:
    wb = openpyxl.load_workbook(str(path))
    ws = wb[marks.SUMMARY_SHEET]
    blue = Font(color=marks.CHECKER_BLUE)
    blue_bold = Font(color=marks.CHECKER_BLUE, bold=True)
    row = ws.max_row + 1
    ws.cell(row=row, column=1, value=marks.REVIEW_PASSED if passed else marks.REVIEW_FAILED).font = blue_bold
    if not passed:
        row += 1
        for col, name in enumerate(FIELDS, 1):
            ws.cell(row=row, column=col, value=name).font = blue_bold
        for d in items:
            row += 1
            for col, name in enumerate(FIELDS, 1):
                ws.cell(row=row, column=col, value=str(d.get(name, ""))).font = blue
        row += 1
        ws.cell(row=row, column=1, value=summary).font = blue
    # openpyxl drops drawing shapes (a student's text boxes) on save: keep a
    # copy of the graded file as it was and put its shapes back afterwards.
    before = path.with_name(path.name + ".before-review.tmp")
    shutil.copy2(path, before)
    try:
        wb.save(str(path))
        restore_shapes(before, path)
    finally:
        before.unlink(missing_ok=True)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("graded_file", type=Path)
    parser.add_argument("result", choices=["passed", "failed"])
    parser.add_argument("discrepancies", type=Path, nargs="?")
    parser.add_argument("--submission", type=Path, required=True,
                        help="the original submission (locates the checker's audit report)")
    args = parser.parse_args()

    try:
        for p in (args.graded_file, args.submission):
            if not p.exists():
                raise ReviewError(f"file not found: {p}")
        ext = args.graded_file.suffix.lower()
        if ext not in (".docx", ".xlsx"):
            raise ReviewError(f"expected a _Graded.docx or _Graded.xlsx, got {ext}")
        check_state(marks.mark_state(args.graded_file))
        passed = args.result == "passed"
        if passed and args.discrepancies:
            raise ReviewError("'passed' takes no discrepancies file")
        if not passed and not args.discrepancies:
            raise ReviewError("'failed' needs the discrepancies JSON")
        items, summary = ([], "") if passed else load_discrepancies(args.discrepancies)
        check_audit(args.graded_file, args.submission, passed, items)
        (mark_document if ext == ".docx" else mark_workbook)(args.graded_file, passed, items, summary)
    except ReviewError as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False))
        sys.exit(1)
    except Exception as exc:  # corrupt graded file, file locked open in Word/Excel, …
        print(json.dumps({"error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False))
        sys.exit(1)
    print(json.dumps({"graded_file": str(args.graded_file),
                      "mark": marks.REVIEW_PASSED if passed else marks.REVIEW_FAILED,
                      "discrepancies": len(items)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
