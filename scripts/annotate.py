#!/usr/bin/env python3
"""Write the grader's annotations into `[name]_Graded.docx|xlsx` — the one
command a grader runs instead of hand-writing python-docx/openpyxl code.

The grader supplies only judgment — which answers are wrong and why — as a
small verdicts JSON; everything mechanical that grader.md pins exactly is
done here, identically for every submission:

- documents (.docx/.doc/.pdf -> .docx): each annotation becomes a new
  paragraph immediately below its anchor, red (FF0000), verdict label bold,
  in the form "INCORRECT: explanation"; `Grading Completed` (red, bold) is
  appended at the end. A .doc is annotated from its Word/LibreOffice
  conversion and a .pdf from the base .docx rebuilt from it
  (lib/pdf_docx.py) — the same documents whose ids the extraction view shows.
- workbooks (.xlsx/.xls -> .xlsx): graded cell by cell. Each annotated
  cell gets Excel's "Bad" style (FFC7CE fill, 9C0006 font) with its
  value/formula untouched, and the explanation goes into the closest empty,
  visible cell — right, then below, then outward along the row, then down
  the column — in bold red (FF0000); a final "Grading Summary" tab holds
  `Grading Completed` in A1. The student's text boxes/shapes (which openpyxl
  drops) are restored. Against a workbook solution the cell rule of
  compare_xlsx.py is enforced: every cell it lists under MARK must be
  annotated with its label (a wrong formula or wrong typed value INCORRECT,
  an empty cell INCOMPLETE), and no correct or carried-over cell (right
  formula, value wrong only because of an upstream error) may be — the
  script refuses verdicts that break the rule.

Provenance (lib/marks.py): the graded file records which solutions file and
which verified rubric it was graded against, and every annotation carries a
hidden tag naming its rubric question — so the checker's audit can map
annotations to questions exactly, and outdated grading can be found later
(`rubric.py graded`). The rubric must be verified; question ids must be the
rubric's. For a workbook `question` may be left out: the cell's rubric
question is recorded.

verdicts.json (write it in the submission's work_dir from extract.py; it is
deleted once the graded file is written — the graded file is the record):
    {"annotations": [
        {"question": "Q3", "verdict": "INCORRECT", "anchor": "p61",
         "text": "The correct answer is C ($300.29) because ... Your answer A ..."},
        {"question": "Q7", "verdict": "INCOMPLETE", "anchor": "t2:r3c1", "text": "..."},
        {"verdict": "INCORRECT", "anchor": "Ex TN5 WACC!G12", "text": "..."}
    ]}
  question: a rubric question id (`rubric.py path` lists them); optional
  for a workbook cell.
  anchor: document `pN` (below paragraph N), `tN` (below table N), `sN`
  (below content control N), or `tN:rRcC` (inside that table cell);
  workbook `Sheet!Cell` (for an answer in a text box: the cell it is
  anchored at, as the view lists it).
  verdict: INCORRECT or INCOMPLETE (a partial answer is INCOMPLETE).
  An empty list is valid (every answer correct) — only the mark is added.

Usage:
    .venv/Scripts/python scripts/annotate.py <submission> <verdicts.json> <output_dir> --solutions <solutions> [--overwrite]
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import warnings
from copy import copy
from pathlib import Path

import _venv  # noqa: F401  — must precede third-party imports (re-runs under .venv)

import docx
import openpyxl
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph
from openpyxl.cell.cell import MergedCell
from openpyxl.chartsheet import Chartsheet
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib import cache, marks, views  # noqa: E402
from lib.xlsx_shapes import restore_shapes  # noqa: E402
import compare_xlsx  # noqa: E402
import extract  # noqa: E402
import rubric  # noqa: E402
from extract_docx import body_blocks  # noqa: E402

DOC_EXTS = {".docx", ".doc", ".pdf"}
SHEET_EXTS = {".xlsx", ".xls"}
BLOCK_ANCHOR_RE = re.compile(r"^[pts]\d+$")
CELL_ANCHOR_RE = re.compile(r"^(t\d+):r(\d+)c(\d+)$")
LABEL_PREFIX_RE = re.compile(r"^\s*\**\s*(INCORRECT|INCOMPLETE|PARTIAL)\s*\**\s*:\s*", re.I)
MAX_SEARCH = 200
BOOKMARK_ID_BASE = 880000


class AnnotateError(Exception):
    pass


def load_annotations(path: Path, question_ids: list[str], question_optional: bool = False) -> list[dict]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise AnnotateError(f"cannot read verdicts JSON: {exc}")
    items = data.get("annotations") if isinstance(data, dict) else None
    if not isinstance(items, list):
        raise AnnotateError('verdicts JSON must be {"annotations": [...]}')
    out = []
    for n, a in enumerate(items):
        where = f"annotations[{n}]"
        if not isinstance(a, dict):
            raise AnnotateError(f"{where}: must be an object")
        if a.get("question") not in question_ids and not (question_optional and a.get("question") is None):
            raise AnnotateError(f"{where}: 'question' must be one of the rubric's question ids "
                                f"{question_ids}, got {a.get('question')!r}")
        verdict = str(a.get("verdict", "")).upper()
        if verdict == "PARTIAL":
            verdict = "INCOMPLETE"
        if verdict not in marks.VERDICT_LABELS:
            raise AnnotateError(f"{where}: verdict must be one of {marks.VERDICT_LABELS} "
                                "(correct answers get no annotation)")
        text = LABEL_PREFIX_RE.sub("", str(a.get("text", ""))).strip()
        if not text:
            raise AnnotateError(f"{where}: 'text' (the explanation) is required")
        mentioned = marks.answer_key_mentions(text)
        if mentioned:
            raise AnnotateError(
                f"{where}: the explanation mentions {mentioned} — students have no access to the "
                "rubric, solution or answer key, so never refer to them. State the correct answer "
                "and working directly (e.g. 'The correct answer is …'); if the word is the "
                "subject's own term, rephrase it")
        anchor = a.get("anchor")
        if not isinstance(anchor, str) or not anchor.strip():
            raise AnnotateError(f"{where}: 'anchor' is required")
        out.append({"question": a.get("question"), "verdict": verdict,
                    "anchor": anchor.strip(), "text": text})
    return out


def cell_id(anchor: str) -> str:
    """A workbook anchor as compare_xlsx.py names cells: Sheet!G12."""
    sheet, ref = views.split_sheet_ref(anchor)
    return f"{sheet}!{ref.replace('$', '').upper()}"


def check_cell_rule(annotations: list[dict], result: dict) -> None:
    """Refuse workbook verdicts that break the cell rule: every MARK cell
    annotated with its label, no correct or carried-over cell annotated. A
    missing question is filled in from the cell's rubric question."""
    must, must_not = compare_xlsx.expected_marks(result)
    question_of = {f"{r['sheet']}!{r['cell']}": r.get("question") for r in result["cells"]}
    labels: dict[str, set] = {}
    for a in annotations:
        if "!" not in a["anchor"]:
            continue  # annotate_workbook reports the bad anchor
        cid = cell_id(a["anchor"])
        labels.setdefault(cid, set()).add(a["verdict"])
        if a["question"] is None:
            a["question"] = question_of.get(cid)
    problems = []
    for cid, r in must.items():
        if cid not in labels:
            problems.append(f"{cid} must be marked {r['mark']}: {compare_xlsx.describe(r)}")
        elif r["mark"] not in labels[cid]:
            problems.append(f"{cid} must be labelled {r['mark']}, not "
                            f"{'/'.join(sorted(labels[cid]))}: {compare_xlsx.describe(r)}")
    for cid, r in must_not.items():
        if cid in labels:
            problems.append(f"{cid} must not be marked: {compare_xlsx.describe(r)}")
    for cid in sorted(set(labels) & set(result.get("not_graded", []))):
        problems.append(f"{cid} must not be marked: the answer key leaves it out of grading "
                        "(a note, label or excluded content)")
    if problems:
        raise AnnotateError("the verdicts break the cell rule (see compare_xlsx.py): "
                            + "; ".join(problems))


def provenance(info: dict, annotations: list[dict], placed: list[dict]) -> dict:
    return {"solutions_file": info["solutions_file"],
            "solutions_sha256": info["solutions_sha256"], "rubric_sha256": info["rubric_sha256"],
            "rubric_round": info["round"],
            "annotations": [{"n": n, "question": a["question"], "anchor": p["anchor"]}
                            for n, (a, p) in enumerate(zip(annotations, placed), 1)]}


# ------------------------------------------------------------------ documents

def _style_annotation(paragraph: Paragraph, verdict: str, text: str, n: int) -> None:
    label = paragraph.add_run(verdict)
    label.bold = True
    body = paragraph.add_run(f": {text}")
    for run in (label, body):
        run.font.color.rgb = RGBColor.from_string(marks.GRADER_RED)
    # hidden tag: a bookmark spanning the annotation, named ga_ann_<n>
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(qn("w:id"), str(BOOKMARK_ID_BASE + n))
    start.set(qn("w:name"), f"{marks.ANNOTATION_TAG}{n}")
    end.set(qn("w:id"), str(BOOKMARK_ID_BASE + n))
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    paragraph._p.append(end)


def _new_paragraph_after(element, parent) -> Paragraph:
    new_p = OxmlElement("w:p")
    element.addnext(new_p)
    return Paragraph(new_p, parent)


def annotate_document(submission: Path, annotations: list[dict], out_path: Path,
                      info: dict) -> list[dict]:
    document = docx.Document(str(extract.source_docx_for(submission)))
    if marks.mark_state_of_document(document)["grading_completed"]:
        raise AnnotateError("source already contains a 'Grading Completed' paragraph")

    # Resolve every anchor against the original numbering before inserting —
    # the same ids the extraction view shows.
    blocks = {bid: obj for bid, _kind, obj in body_blocks(document)}
    last_inserted: dict[str, Paragraph] = {}
    placed = []
    for n, a in enumerate(annotations, 1):
        anchor = a["anchor"]
        m_cell = CELL_ANCHOR_RE.match(anchor)
        if BLOCK_ANCHOR_RE.match(anchor):
            if anchor not in blocks:
                raise AnnotateError(f"anchor {anchor}: no such block in the document")
            prev = last_inserted.get(anchor)
            if prev is not None:
                element = prev._p
            else:
                obj = blocks[anchor]
                element = obj._p if anchor[0] == "p" else obj._tbl if anchor[0] == "t" else obj
            para = _new_paragraph_after(element, document)
        elif m_cell:
            table_id, r, c = m_cell.group(1), int(m_cell.group(2)), int(m_cell.group(3))
            if table_id not in blocks:
                raise AnnotateError(f"anchor {anchor}: no table {table_id}")
            try:
                cell = blocks[table_id].cell(r, c)
            except IndexError:
                raise AnnotateError(f"anchor {anchor}: no such cell in table {table_id}")
            para = cell.add_paragraph()
        else:
            raise AnnotateError(f"anchor {anchor!r}: expected pN, tN, sN or tN:rRcC for a document")
        _style_annotation(para, a["verdict"], a["text"], n)
        last_inserted[anchor] = para
        placed.append({"question": a["question"], "anchor": anchor})

    mark = document.add_paragraph()
    run = mark.add_run(marks.GRADING_COMPLETED)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(marks.GRADER_RED)
    marks.write_meta_docx(document, provenance(info, annotations, placed))
    document.save(str(out_path))
    return placed


# ------------------------------------------------------------------ workbooks

def _workbook_source(submission: Path) -> tuple[Path | None, list[str]]:
    """The .xlsx a workbook's graded copy is written from: the file itself,
    or the .xls's Excel/LibreOffice conversion made during extraction (so
    the view's cell references are exactly the ones annotated)."""
    if submission.suffix.lower() == ".xlsx":
        return submission, []
    record, *_ = extract.get_record(submission)
    if record.get("converter_output_path"):
        source = Path(record["converter_output_path"])
        if not source.exists():  # cache partly cleaned — reconvert
            record, *_ = extract.get_record(submission, force=True)
            source = Path(record["converter_output_path"])
        return source, []
    return None, ["DEGRADED: no .xls converter (Microsoft Excel or LibreOffice) is installed — the "
                  "graded workbook was rebuilt from cached values only (formulas and formatting "
                  "lost). Install Excel or LibreOffice and regrade for a faithful copy."]


def _rebuild_from_xls(submission: Path) -> openpyxl.Workbook:
    record, *_ = extract.get_record(submission)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in record["sheet_order"]:
        ws = wb.create_sheet(name)
        for coord, cell in record["sheets"][name]["cells"].items():
            if cell.get("value") not in (None, ""):
                ws[coord] = cell["value"]
    return wb


def hidden_rows_cols(ws) -> tuple[set[int], set[int]]:
    """Rows and columns hidden on a sheet (column dimensions may span a
    range). Feedback written there would never be seen."""
    rows = {r for r, dim in ws.row_dimensions.items() if dim.hidden}
    cols = set()
    for dim in ws.column_dimensions.values():
        if dim.hidden and dim.min and dim.max:
            cols.update(range(dim.min, dim.max + 1))
    return rows, cols


def search_order(row: int, col: int):
    """Candidate cells for an answer's explanation, closest first: right,
    below, then outward along the row, then down the column. audit_graded.py
    replays the same order."""
    yield row, col + 1
    yield row + 1, col
    for c in range(col + 2, col + MAX_SEARCH):
        yield row, c
    for r in range(row + 2, row + MAX_SEARCH):
        yield r, col


def is_free(ws, row: int, col: int, taken: set, hidden: tuple[set, set]) -> bool:
    if (row, col) in taken or row in hidden[0] or col in hidden[1]:
        return False
    cell = ws.cell(row=row, column=col)
    return not isinstance(cell, MergedCell) and cell.value in (None, "")


def closest_empty(ws, row: int, col: int, taken: set, hidden: tuple[set, set]) -> tuple[int, int]:
    for r, c in search_order(row, col):
        if is_free(ws, r, c, taken, hidden):
            return r, c
    raise AnnotateError(f"no empty visible cell near {get_column_letter(col)}{row}")


def annotate_workbook(submission: Path, annotations: list[dict], out_path: Path,
                      info: dict) -> tuple[list, list]:
    source, notes = _workbook_source(submission)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(str(source)) if source else _rebuild_from_xls(submission)
    if marks.SUMMARY_SHEET in wb.sheetnames:
        raise AnnotateError(f"workbook already has a '{marks.SUMMARY_SHEET}' tab")

    red_bold = Font(color=marks.GRADER_RED, bold=True)
    bad_fill = PatternFill(start_color=marks.BAD_FILL, end_color=marks.BAD_FILL, fill_type="solid")
    targets = []
    for n, a in enumerate(annotations, 1):
        if "!" not in a["anchor"]:
            raise AnnotateError(f"anchor {a['anchor']!r}: expected Sheet!Cell for a workbook")
        sheet, ref = views.split_sheet_ref(a["anchor"])
        if sheet not in wb.sheetnames:
            raise AnnotateError(f"anchor {a['anchor']!r}: no sheet {sheet!r}")
        ws = wb[sheet]
        if isinstance(ws, Chartsheet):
            raise AnnotateError(f"anchor {a['anchor']!r}: {sheet!r} is a chart sheet (no cells) — "
                                "anchor at the chart's data cells on another tab")
        try:
            row, col = coordinate_to_tuple(ref.replace("$", "").upper())
        except ValueError:
            raise AnnotateError(f"anchor {a['anchor']!r}: bad cell reference")
        for rng in ws.merged_cells.ranges:  # a merged answer lives in its top-left cell
            if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
                row, col = rng.min_row, rng.min_col
        targets.append((n, a, sheet, row, col))

    # Every answer cell is reserved before any explanation is placed (an
    # empty answer cell must never receive another cell's explanation), and
    # explanations are placed answer by answer in row/column order, several
    # on one answer stacking in list order — exactly the order
    # audit_graded.py replays.
    taken: dict[str, set] = {}
    hidden: dict[str, tuple] = {}
    for _n, _a, sheet, row, col in targets:
        taken.setdefault(sheet, set()).add((row, col))
    placed: list = [None] * len(annotations)
    for n, a, sheet, row, col in sorted(targets, key=lambda t: (t[2], t[3], t[4], t[0])):
        ws = wb[sheet]
        sheet_taken = taken[sheet]
        sheet_hidden = hidden.setdefault(sheet, hidden_rows_cols(ws))

        answer = ws.cell(row=row, column=col)
        answer.fill = bad_fill
        font = copy(answer.font)
        font.color = marks.BAD_FONT
        answer.font = font

        r, c = closest_empty(ws, row, col, sheet_taken, sheet_hidden)
        sheet_taken.add((r, c))
        note = ws.cell(row=r, column=c)
        note.value = f"{a['verdict']}: {a['text']}"
        note.font = red_bold
        tag = f"{marks.ANNOTATION_TAG}{n}"
        wb.defined_names[tag] = DefinedName(
            tag, attr_text=f"{quote_sheetname(sheet)}!${get_column_letter(c)}${r}")
        placed[n - 1] = {"question": a["question"], "anchor": f"{sheet}!{answer.coordinate}",
                         "annotation_cell": f"{sheet}!{note.coordinate}"}

    summary = wb.create_sheet(marks.SUMMARY_SHEET)
    summary["A1"] = marks.GRADING_COMPLETED
    summary["A1"].font = red_bold
    marks.write_meta_xlsx(wb, provenance(info, annotations, placed))
    wb.save(str(out_path))
    if source is not None:
        restored = restore_shapes(source, out_path)
        if restored:
            notes.append(f"restored {restored} text box(es)/shape(s) openpyxl does not keep")
    return placed, notes


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("submission", type=Path)
    parser.add_argument("verdicts", type=Path)
    parser.add_argument("output_dir", type=Path)
    parser.add_argument("--solutions", type=Path, required=True,
                        help="the matched solutions file — its verified rubric is recorded")
    parser.add_argument("--overwrite", action="store_true",
                        help="replace your own existing, unreviewed _Graded file (a reviewed "
                             "one is never overwritten)")
    args = parser.parse_args()

    ext = args.submission.suffix.lower()
    for p in (args.submission, args.solutions):
        if not p.exists():
            print(json.dumps({"error": f"file not found: {p}"}))
            sys.exit(1)
    if ext not in DOC_EXTS | SHEET_EXTS:
        print(json.dumps({"error": f"unsupported file type: {ext}"}))
        sys.exit(1)
    out_ext = ".xlsx" if ext in SHEET_EXTS else ".docx"
    args.output_dir.mkdir(parents=True, exist_ok=True)
    out_path = args.output_dir / f"{args.submission.stem}_Graded{out_ext}"
    try:
        info = rubric.verified_rubric(args.solutions)
        if out_path.exists():
            if not args.overwrite:
                raise AnnotateError(f"{out_path} already exists — refusing to overwrite (pass "
                                    "--overwrite only for your own stale, unreviewed output)")
            if marks.mark_state(out_path)["reviewed"]:
                raise AnnotateError(f"{out_path} already carries the checker's review — it is "
                                    "never overwritten")
        cell_mode = ext in SHEET_EXTS and args.solutions.suffix.lower() in SHEET_EXTS
        annotations = load_annotations(args.verdicts, info["question_ids"], question_optional=cell_mode)
        if cell_mode:
            check_cell_rule(annotations, compare_xlsx.grade_cells(args.submission, args.solutions))
        if ext in SHEET_EXTS:
            placed, notes = annotate_workbook(args.submission, annotations, out_path, info)
        else:
            placed, notes = annotate_document(args.submission, annotations, out_path, info), []
    except (AnnotateError, rubric.RubricError) as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False))
        sys.exit(1)
    except Exception as exc:  # conversion failure, corrupt file, …
        print(json.dumps({"error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False))
        sys.exit(1)

    # The graded file now holds the verdicts; don't leave them lying in the
    # shared scratch folder for the checker to see before it grades blind.
    try:
        if args.verdicts.resolve().is_relative_to(cache.WORK_DIR.resolve()):
            args.verdicts.unlink()
    except OSError:
        pass
    print(json.dumps({"graded_file": str(out_path), "annotations": len(placed),
                      "placed": placed, "notes": notes, "rubric_round": info["round"],
                      "mark": marks.GRADING_COMPLETED}, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
