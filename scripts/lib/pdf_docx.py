"""Rebuild a PDF as a plain .docx, deterministically.

A PDF submission is graded into a `_Graded.docx`, which previously meant the
grader retyping the whole document through python-docx by hand — expensive in
output tokens and a chance to silently drop or alter student content. This
builds that base document once, in code, from PyMuPDF's layout analysis:

- each text line becomes one paragraph (bold/italic spans kept) — PyMuPDF
  text blocks often span several questions, and a line is the finest unit an
  annotation can then be placed directly below;
- equations whose text layer is corrupted (Word-exported math comes out as
  doubled Mathematical Alphanumeric Symbols, e.g. "𝑃𝑃𝑃𝑃=") are re-inserted
  as a picture of their on-page region instead, so the graded copy shows the
  student's equation exactly as written rather than garbage text. An
  equation's region is its math lines plus plain lines that belong to it
  (a bare numerator/denominator between math lines, or text in the same
  vertical band beside it);
- each embedded raster image is re-inserted as a picture at its position in
  reading order (rendered from its on-page region, so masks/odd encodings
  come out right);
- vector graphics — a chart exported from Excel, a diagram drawn with shapes
  — are re-inserted as a picture of their region too (text lines inside the
  region, e.g. axis labels, belong to the picture), so they are neither lost
  from the graded copy nor shredded into stray one-word paragraphs;
- the first paragraph of every page after the first starts a new page.

extract.py extracts the result like any .docx, so its `[pN]` ids are exactly
the paragraph indices annotate.py inserts annotations after.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

import docx
import pymupdf
from docx.shared import Inches, Pt

MAX_WIDTH_IN = 6.5
BOLD_FLAG = 16
ITALIC_FLAG = 2
MATH_ALPHANUMERIC_RE = re.compile("[\U0001D400-\U0001D7FF]")
BAND_TOLERANCE = 2.0  # points
# A vector region counts as a graphic only if it is at least this big —
# smaller clusters are highlights, underlines, bullets, checkbox squares.
MIN_GRAPHIC_SIDE = 40.0
MIN_GRAPHIC_AREA = 3000.0
GRAY_GRAPHIC_PATHS = 8


def _colored(color) -> bool:
    return bool(color) and len(color) >= 3 and max(color) - min(color) >= 0.15


LABEL_MARGIN = 24.0  # points around a graphic searched for its labels
LABEL_MAX_CHARS = 40


def _with_labels(page, rect):
    """Grow a graphic's region over the short text lines hugging it — a
    chart's title, axis numbers, legend — so they are pictured with it
    instead of becoming stray one-word paragraphs."""
    lines = [pymupdf.Rect(line["bbox"])
             for block in page.get_text("dict")["blocks"] if block.get("type") == 0
             for line in block["lines"]
             if 0 < len(_line_text(line).strip()) <= LABEL_MAX_CHARS]
    region = pymupdf.Rect(rect)
    grown = True
    while grown:
        grown = False
        halo = pymupdf.Rect(region.x0 - LABEL_MARGIN, region.y0 - LABEL_MARGIN,
                            region.x1 + LABEL_MARGIN, region.y1 + LABEL_MARGIN)
        for line in lines:
            if line.intersects(halo) and not line in region:
                region |= line
                grown = True
    return region & page.rect


def graphic_regions(page) -> tuple[list, bool]:
    """(regions to picture, whether the page holds vector graphics at all).

    A large drawing cluster with colour or curves (a chart, a diagram) is
    pictured. A large greyscale cluster with many paths may be a chart — or
    just a bordered table — so it only flags the page for visual reading and
    its text stays text."""
    drawings = page.get_drawings()
    if not drawings:
        return [], False
    pictured, flagged = [], False
    for rect in page.cluster_drawings(drawings=drawings):
        if (rect.width < MIN_GRAPHIC_SIDE or rect.height < MIN_GRAPHIC_SIDE
                or rect.width * rect.height < MIN_GRAPHIC_AREA):
            continue
        inside = [d for d in drawings if pymupdf.Rect(d["rect"]).intersects(rect)]
        if any(_colored(d.get("fill")) or _colored(d.get("color"))
               or any(item[0] == "c" for item in d["items"]) for d in inside):
            pictured.append(_with_labels(page, rect))
            flagged = True
        elif sum(1 for d in inside if d.get("fill") not in (None, (1.0, 1.0, 1.0))
                 or d.get("color")) >= GRAY_GRAPHIC_PATHS:
            flagged = True
    return pictured, flagged


def _line_text(line) -> str:
    return "".join(span["text"] for span in line["spans"])


def _page_items(page) -> list[dict]:
    """Image blocks, vector-graphic regions and non-blank text lines, in
    reading order."""
    graphics, _ = graphic_regions(page)

    def in_graphic(rect) -> bool:
        center = pymupdf.Point((rect.x0 + rect.x1) / 2, (rect.y0 + rect.y1) / 2)
        return any(center in g for g in graphics)

    items = []
    for block in page.get_text("dict", sort=True)["blocks"]:
        if block.get("type") == 1:
            if not in_graphic(pymupdf.Rect(block["bbox"])):
                items.append({"kind": "image", "bbox": pymupdf.Rect(block["bbox"])})
            continue
        first = True
        for line in block.get("lines", []):
            text = _line_text(line)
            if not text.strip() or in_graphic(pymupdf.Rect(line["bbox"])):
                continue
            items.append({"kind": "line", "line": line, "bbox": pymupdf.Rect(line["bbox"]),
                          "math": bool(MATH_ALPHANUMERIC_RE.search(text)),
                          "block_start": first})
            first = False
    for g in graphics:  # each graphic goes in before the first item below its top
        at = next((i for i, it in enumerate(items) if it["bbox"].y0 >= g.y0), len(items))
        items.insert(at, {"kind": "image", "bbox": g})
    return items


def _equation_clusters(items: list[dict]) -> dict[int, int]:
    """first item index -> last item index of each equation region."""
    clusters: dict[int, int] = {}
    last_start = None
    for i, item in enumerate(items):
        if not item.get("math"):
            continue
        if last_start is not None and i - clusters[last_start] <= 2 and all(
                it["kind"] == "line" for it in items[clusters[last_start]:i + 1]):
            clusters[last_start] = i  # at most one plain line in between
        else:
            clusters[i] = i
            last_start = i
    # Grow each region over adjacent plain lines in its vertical band
    # (e.g. the rest of an equation laid out beside the math glyphs).
    grown: dict[int, int] = {}
    taken_until = -1
    for start in sorted(clusters):
        end = clusters[start]
        start = max(start, taken_until + 1)
        band = pymupdf.Rect(items[start]["bbox"])
        for it in items[start:end + 1]:
            band |= it["bbox"]

        def in_band(it) -> bool:
            # Same vertical band and not left of the equation — a label such
            # as "Answer: B" sitting to its left stays readable text.
            mid = (it["bbox"].y0 + it["bbox"].y1) / 2
            return (it["kind"] == "line"
                    and band.y0 - BAND_TOLERANCE <= mid <= band.y1 + BAND_TOLERANCE
                    and it["bbox"].x0 >= band.x0 - BAND_TOLERANCE)

        while start - 1 > taken_until and in_band(items[start - 1]):
            start -= 1
            band |= items[start]["bbox"]
        while end + 1 < len(items) and in_band(items[end + 1]):
            end += 1
            band |= items[end]["bbox"]
        grown[start] = end
        taken_until = end
    return grown


def _add_region_picture(document, page, rect):
    rect = pymupdf.Rect(rect) & page.rect
    if rect.is_empty or rect.width < 2 or rect.height < 2:
        return None
    pix = page.get_pixmap(clip=rect, dpi=150)
    paragraph = document.add_paragraph()
    paragraph.add_run().add_picture(
        io.BytesIO(pix.tobytes("png")),
        width=Inches(min(rect.width / 72, MAX_WIDTH_IN)),
    )
    return paragraph


def _add_text_line(document, item):
    paragraph = document.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(0)
    fmt.space_before = Pt(6) if item["block_start"] else Pt(0)
    for span in item["line"]["spans"]:
        if not span["text"]:
            continue
        run = paragraph.add_run(span["text"])
        if span["flags"] & BOLD_FLAG:
            run.bold = True
        if span["flags"] & ITALIC_FLAG:
            run.italic = True
    return paragraph


def build_base_docx(pdf_path: Path, out_path: Path) -> list[int]:
    """Write the rebuilt document to `out_path`; return the 0-based page
    index of each top-level paragraph, in order."""
    src = pymupdf.open(str(pdf_path))
    document = docx.Document()
    page_of_paragraph: list[int] = []

    for page_index, page in enumerate(src):
        items = _page_items(page)
        clusters = _equation_clusters(items)
        new_paragraphs = []
        i = 0
        while i < len(items):
            item = items[i]
            if i in clusters:
                end = clusters[i]
                region = pymupdf.Rect(item["bbox"])
                for it in items[i + 1:end + 1]:
                    region |= it["bbox"]
                paragraph = _add_region_picture(document, page, region)
                i = end + 1
            elif item["kind"] == "image":
                paragraph = _add_region_picture(document, page, item["bbox"])
                i += 1
            else:
                paragraph = _add_text_line(document, item)
                i += 1
            if paragraph is not None:
                new_paragraphs.append(paragraph)

        if new_paragraphs and page_index > 0:
            new_paragraphs[0].paragraph_format.page_break_before = True
        page_of_paragraph += [page_index] * len(new_paragraphs)

    src.close()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))
    return page_of_paragraph
